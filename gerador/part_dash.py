"""DASHBOARD (indicadores, saúde financeira, alertas, comparativos, metas e 8 gráficos), COMEÇAR AQUI e BACKUP."""
from openpyxl.chart import BarChart, LineChart, DoughnutChart, Reference, Series
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.title import Title
from openpyxl.chart.text import RichText, Text
from openpyxl.drawing.text import RichTextProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, RegularTextRun
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.pagebreak import Break
from common import *
from part_base import N_CATD, N_METAS, N_CART, MESES_AB
from part_cards_debts import N_DIV


# ---------------------------------------------------------------- indicadores e alertas
def build_kpis(ctx):
    c = ctx.calc
    c.sec('Indicadores do mês selecionado')
    for nome, tipo in (('Rec', 'Receita'), ('Desp', 'Despesa'), ('Apo', 'Investimento')):
        c.val(f'K_{nome}', f'={SM(tipo, "Sel_Idx")}', f'{tipo}s do mês', BRLZ)
        c.val(f'K_{nome}Ant', f'={SM(tipo, "Sel_Idx-1")}', f'{tipo}s do mês anterior', BRLZ)
        c.val(f'K_{nome}AA', f'={SM(tipo, "Sel_Idx-12")}', f'{tipo}s do mesmo mês do ano anterior', BRLZ)
    c.val('K_Saldo', '=K_Rec-K_Desp', 'Saldo do mês', BRLZ)
    c.val('K_SaldoAnt', '=K_RecAnt-K_DespAnt', 'Saldo do mês anterior', BRLZ)
    c.val('K_SaldoAA', '=K_RecAA-K_DespAA', 'Saldo do mesmo mês do ano anterior', BRLZ)
    c.val('K_Poup', '=IFERROR(K_Saldo/K_Rec,0)', 'Taxa de poupança', PCT)
    c.val('K_PoupAnt', '=IFERROR(K_SaldoAnt/K_RecAnt,0)', 'Taxa de poupança (mês anterior)', PCT)
    c.val('K_PoupAA', '=IFERROR(K_SaldoAA/K_RecAA,0)', 'Taxa de poupança (ano anterior)', PCT)
    c.val('K_PoupMeta', '=IF(N(Cfg_MetaPoup)=0,0.2,Cfg_MetaPoup)', 'Meta de poupança efetiva', PCT)
    c.val('K_SaldoContas', '=SUM(Contas_SaldoAtual)', 'Saldo total em contas', BRLZ)
    c.val('K_Reserva', '=SUMIFS(Meta_Guardado,Meta_Cat,"Reserva")', 'Reserva de emergência (metas da categoria Reserva)', BRLZ)
    c.val('K_CustoMes', '=IF(K_Med3Desp>0,K_Med3Desp,K_RendaBase)', 'Custo mensal de referência', BRLZ)
    c.val('K_ReservaMeses', '=IFERROR(K_Reserva/K_CustoMes,0)', 'Meses cobertos pela reserva', NUM1)
    c.val('K_MesesAlvo', '=IF(N(Cfg_MesesReserva)=0,6,Cfg_MesesReserva)', 'Meses-alvo da reserva', INT)
    c.val('K_ReservaPct', '=IFERROR(K_Reserva/(K_MesesAlvo*K_CustoMes),0)', '% da reserva-alvo', PCT)
    c.val('K_GastoDia', '=IFERROR(K_Desp/IF(Sel_Idx=Cfg_Idx,MAX(1,DAY(Cfg_Hoje)),Sel_Dias),0)', 'Gasto médio diário', BRLZ)
    c.val('K_DiasCaixa', '=IF(K_GastoDia>0,MAX(0,K_SaldoContas)/K_GastoDia,0)', 'Dias de caixa restantes', INT)
    c.val('K_AporteMeta', '=IF(N(Cfg_MetaAporte)>0,Cfg_MetaAporte,K_PoupMeta*K_RendaBase)', 'Poupança/aporte planejado do mês', BRLZ)
    c.val('K_AportePct', '=IFERROR(K_Apo/K_AporteMeta,0)', '% do aporte planejado', PCT)
    c.sec('Nota de saúde financeira (0 a 100)')
    c.val('H_Poup', '=20*MIN(1,MAX(0,IFERROR(K_Poup/K_PoupMeta,0)))', 'Taxa de poupança (máx. 20)', NUM1)
    c.val('H_Div', '=20*MAX(0,MIN(1,(0.4-Div_PctParc)/0.3))', 'Endividamento sobre a renda (máx. 20)', NUM1)
    c.val('H_Res', '=20*MIN(1,IFERROR(K_ReservaMeses/K_MesesAlvo,0))', 'Reserva de emergência (máx. 20)', NUM1)
    c.val('H_Orc', '=IF(Orc_QtdPlan=0,10,20*SUMPRODUCT((Orc_Plan>0)*(Orc_Real<=Orc_Plan))/Orc_QtdPlan)',
          'Gastos dentro do orçamento (máx. 20)', NUM1)
    c.val('H_Apo', '=20*Conq_Apo6/6', 'Regularidade de aportes (máx. 20)', NUM1)
    c.val('H_Nota', '=ROUND(H_Poup+H_Div+H_Res+H_Orc+H_Apo,0)', 'NOTA FINAL', INT)
    c.val('H_Sem', '=IF(H_Nota>=70,"Saudável",IF(H_Nota>=40,"Atenção","Crítico"))', 'Semáforo')

    # ---- alertas
    ws, sh = c.ws, c.sh
    sh.put('AH5', 'Alertas automáticos (candidatos)', 'sec')
    sh.put('AH6', 'Mensagem', 'hdr')
    sh.put('AI6', 'Contagem', 'hdr')
    ws.column_dimensions['AH'].width = 70
    cand = []
    cand.append('=IF(AND(K_Rec+K_Desp>0,K_Saldo<0),"🔴 Saldo do mês negativo: despesas superam as receitas em R$ "&FIXED(-K_Saldo,0)&".","")')
    cand.append('=IF(Cal_Atraso>0,"🔴 Você tem R$ "&FIXED(Cal_Atraso,0)&" em contas atrasadas (veja CALENDÁRIO DE CONTAS).","")')
    cand.append('=IF(AND(K_Med3Desp>0,K_ReservaMeses<K_MesesAlvo),"🟡 Reserva cobre apenas "&FIXED(K_ReservaMeses,1)&" meses (meta: "&K_MesesAlvo&").","")')
    cand.append('=IF(AND(Div_ParcTot>0,Div_PctParc>Div_Lim),"🔴 Parcelas de dívidas = "&ROUND(Div_PctParc*100,0)&"% da renda (limite: "&ROUND(Div_Lim*100,0)&"%).","")')
    for i in range(1, N_METAS + 1):
        cand.append(f'=IF(INDEX(Meta_Status,{i})="Atrasada","🔴 Meta "&INDEX(Meta_Nome,{i})&" está atrasada: faltam R$ "&'
                    f'FIXED(MAX(0,INDEX(Meta_Nec,{i})-INDEX(Meta_Ritmo,{i})),0)&" por mês para recuperar.","")')
    for i in range(1, N_CATD + 1):
        cand.append(f'=IF(AND(INDEX(Orc_Plan,{i})>0,INDEX(Orc_Pct,{i})>=1),"🔴 Orçamento de "&INDEX(Orc_Cat,{i})&" estourado: "&ROUND(INDEX(Orc_Pct,{i})*100,0)&"% usado.",'
                    f'IF(AND(INDEX(Orc_Plan,{i})>0,INDEX(Orc_Pct,{i})>=0.9),"🟡 Você já usou "&ROUND(INDEX(Orc_Pct,{i})*100,0)&"% do orçamento de "&INDEX(Orc_Cat,{i})&".",""))')
    for i in range(1, N_CATD + 1):
        cand.append(f'=IF(AND(INDEX(Orc_Med3,{i})>0,INDEX(Orc_Real,{i})>1.2*INDEX(Orc_Med3,{i})),"🟡 Gasto com "&INDEX(Orc_Cat,{i})&" "&'
                    f'ROUND((INDEX(Orc_Real,{i})/INDEX(Orc_Med3,{i})-1)*100,0)&"% acima da média dos últimos 3 meses.","")')
    for k in range(1, N_CART + 1):
        cand.append(f'=IF(INDEX(Card_Nome,{k})="","",IF(AND(INDEX(Card_FatAnt,{k})>0,N(INDEX(Card_VencAnt,{k}))-Cfg_Hoje>=0,N(INDEX(Card_VencAnt,{k}))-Cfg_Hoje<=7),'
                    f'"🟡 Fatura do cartão "&INDEX(Card_Nome,{k})&" vence em "&(INDEX(Card_VencAnt,{k})-Cfg_Hoje)&" dia(s): R$ "&FIXED(INDEX(Card_FatAnt,{k}),0)&".",""))')
        cand.append(f'=IF(INDEX(Card_Nome,{k})="","",IF(INDEX(Card_Pct,{k})>Cfg_LimCartao,"🟡 Cartão "&INDEX(Card_Nome,{k})&" com "&ROUND(INDEX(Card_Pct,{k})*100,0)&"% do limite usado.",""))')
    cand.append('=IF(COUNTIF(L_Alerta,"⚠ Possível*")>0,"🟡 "&COUNTIF(L_Alerta,"⚠ Possível*")&" lançamento(s) com possível duplicidade (veja LANÇAMENTOS).","")')
    cand.append('=IF(COUNTIF(L_Alerta,"⚠ Categoria*")>0,"🟡 "&COUNTIF(L_Alerta,"⚠ Categoria*")&" lançamento(s) com categoria incompatível com o tipo.","")')
    n = len(cand)
    for i, f in enumerate(cand):
        r = 7 + i
        sh.put(f'AH{r}', f, 'calc')
        sh.put(f'AI{r}', f'=N(AI{r - 1})+IF(AH{r}<>"",1,0)' if i else f'=IF(AH{r}<>"",1,0)', 'calc', INT)
    ctx.name('Alert_Msg', 'CÁLCULOS', f'AH7:AH{6 + n}')
    ctx.name('Alert_Cnt', 'CÁLCULOS', f'AI7:AI{6 + n}')
    ctx.name_formula('Alert_Total', f"MAX('CÁLCULOS'!$AI$7:$AI${6 + n})")


# ---------------------------------------------------------------- gráficos
def rref(ctx, nome):
    sheet, ref = ctx.refs[nome]
    c1, r1, c2, r2 = range_boundaries(ref)
    return Reference(ctx.wb[sheet], min_col=c1, min_row=r1, max_col=c2, max_row=r2)


def _rich(sz, bold=False):
    cp = CharacterProperties(sz=sz, b=bold, solidFill=COR['texto'])
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def make_title(txt, sz=1100):
    cp = CharacterProperties(sz=sz, b=True, solidFill=COR['texto'])
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=txt)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


CW, CH = 14.2, 7.6


def style_chart(ch, titulo, w=CW, h=CH, legend='b', axes=True):
    ch.title = make_title(titulo)
    ch.width, ch.height = w, h
    ch.visible_cells_only = False
    if legend:
        ch.legend.position = legend
        ch.legend.txPr = _rich(800)
    else:
        ch.legend = None
    if axes:
        for ax in (ch.x_axis, ch.y_axis):
            ax.delete = False
            ax.txPr = _rich(800)
        ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E4E7EC'))


def color_series(s, hexc, line=False):
    if line:
        s.graphicalProperties.line.solidFill = hexc
        s.graphicalProperties.line.width = 28000
        s.smooth = False
    else:
        s.graphicalProperties.solidFill = hexc
        s.graphicalProperties.line.solidFill = hexc


def ser(ctx, nome, titulo):
    return Series(rref(ctx, nome), title=titulo)


PALETA_ROSCA = ['2F6FD0', '2E9E5B', 'F2A93B', '7B5EA7', 'D64545', '1ABC9C', '8E9AAF', 'E67E22', 'A0A7B0']


def rosca(ctx, vals, cats, titulo):
    ch = DoughnutChart()
    ch.holeSize = 55
    style_chart(ch, titulo, legend='r', axes=False)
    ch.add_data(rref(ctx, vals), titles_from_data=False)
    ch.set_categories(rref(ctx, cats))
    s = ch.series[0]
    for i, col in enumerate(PALETA_ROSCA):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = col
        s.dPt.append(pt)
    s.dLbls = DataLabelList()
    s.dLbls.showPercent = True
    s.dLbls.showVal = False
    s.dLbls.showCatName = False
    s.dLbls.showSerName = False
    s.dLbls.showLegendKey = False
    s.dLbls.txPr = _rich(800, True)
    return ch


def build_charts(ctx, ws, top):
    # (a) colunas: receitas x despesas
    a = BarChart()
    a.type = 'col'
    a.grouping = 'clustered'
    style_chart(a, 'Receitas x despesas por mês (ano selecionado)')
    s1, s2 = ser(ctx, 'Flux_Rec', 'Receitas'), ser(ctx, 'Flux_Desp', 'Despesas')
    color_series(s1, COR['verde'])
    color_series(s2, COR['vermelho'])
    a.series += [s1, s2]
    a.set_categories(rref(ctx, 'Flux_Mes'))
    a.y_axis.number_format = '#,##0'
    a.gapWidth = 60
    ws.add_chart(a, f'B{top}')
    # (b) rosca despesas por categoria
    ws.add_chart(rosca(ctx, 'Rk_Val', 'Rk_Nome', 'Despesas por categoria no mês'), f'G{top}')
    # (c) linha saldo acumulado (eixo esq.) e patrimônio (eixo dir.)
    c = LineChart()
    style_chart(c, 'Saldo acumulado e patrimônio no ano')
    s1 = ser(ctx, 'Flux_Acum', 'Saldo acumulado (eixo esq.)')
    color_series(s1, COR['verde'], True)
    c.series.append(s1)
    c.set_categories(rref(ctx, 'Flux_Mes'))
    c.y_axis.number_format = '#,##0'
    c2 = LineChart()
    s2 = ser(ctx, 'Pat_Serie', 'Patrimônio líquido (eixo dir.)')
    color_series(s2, COR['azul'], True)
    c2.series.append(s2)
    c2.y_axis.axId = 200
    c2.y_axis.crosses = 'max'
    c2.y_axis.number_format = '#,##0'
    c2.y_axis.delete = False
    c2.y_axis.txPr = _rich(800)
    c2.y_axis.majorGridlines = None
    c2.visible_cells_only = False
    c += c2
    ws.add_chart(c, f'B{top + 16}')
    # (d) barras orçado x realizado
    d = BarChart()
    d.type = 'bar'
    d.grouping = 'clustered'
    style_chart(d, 'Orçado x realizado por categoria (mês)', h=CH + 1.2)
    s1, s2 = ser(ctx, 'OR_Orc', 'Orçado'), ser(ctx, 'OR_Real', 'Realizado')
    color_series(s1, COR['azul'])
    color_series(s2, COR['vermelho'])
    d.series += [s1, s2]
    d.set_categories(rref(ctx, 'OR_Nome'))
    d.x_axis.scaling.orientation = 'maxMin'
    d.x_axis.tickLblSkip = 1
    d.y_axis.crosses = 'max'
    d.y_axis.number_format = '#,##0'
    d.gapWidth = 40
    ws.add_chart(d, f'G{top + 16}')
    # (e) rosca investimentos
    ws.add_chart(rosca(ctx, 'Inv_TipoVal', 'Inv_TipoNome', 'Composição dos investimentos por tipo'), f'B{top + 34}')
    # (f) barras de progresso das metas
    f = BarChart()
    f.type = 'bar'
    f.grouping = 'clustered'
    style_chart(f, 'Progresso das metas (% concluído)', legend=None, h=CH + 1.2)
    s1 = ser(ctx, 'Meta_Pct', '% concluído')
    color_series(s1, COR['roxo'])
    f.series.append(s1)
    f.set_categories(rref(ctx, 'Meta_Nome'))
    f.x_axis.scaling.orientation = 'maxMin'
    f.x_axis.tickLblSkip = 1
    f.y_axis.crosses = 'max'
    f.y_axis.scaling.min = 0
    f.y_axis.scaling.max = 1
    f.y_axis.number_format = '0%'
    f.gapWidth = 40
    ws.add_chart(f, f'G{top + 34}')
    # (g) colunas regra 50/30/20
    g = BarChart()
    g.type = 'col'
    g.grouping = 'clustered'
    style_chart(g, 'Regra 50/30/20: ideal x real (% da receita)')
    s1, s2 = ser(ctx, 'R50_Ideal', 'Ideal'), ser(ctx, 'R50_Real', 'Real')
    color_series(s1, COR['azul'])
    color_series(s2, COR['amarelo'])
    g.series += [s1, s2]
    g.set_categories(rref(ctx, 'R50_Grupo'))
    g.y_axis.number_format = '0%'
    g.gapWidth = 60
    ws.add_chart(g, f'B{top + 52}')
    # (h) linha projeção de saldo 12 meses
    h = LineChart()
    style_chart(h, 'Projeção do saldo acumulado: próximos 12 meses')
    for key, nm, col in (('Sc1_Acum', 'Realista', COR['azul']), ('Sc2_Acum', 'Otimista', COR['verde']),
                         ('Sc3_Acum', 'Pessimista', COR['vermelho'])):
        s = ser(ctx, key, nm)
        color_series(s, col, True)
        h.series.append(s)
    h.set_categories(rref(ctx, 'Proj_Labels'))
    h.y_axis.number_format = '#,##0'
    h.x_axis.txPr = RichText(bodyPr=RichTextProperties(rot=-2700000), p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=700)), endParaRPr=CharacterProperties(sz=700))])
    ws.add_chart(h, f'G{top + 52}')


# ---------------------------------------------------------------- dashboard
def build_dashboard(ctx):
    sh = Sh(ctx, 'DASHBOARD')
    ws = sh.ws
    S = 'DASHBOARD'
    sh.header('DASHBOARD', 'Visão geral do mês. Escolha o mês e o ano abaixo (células azuis) e tudo se atualiza. '
              'Para digitar dados, use a aba LANÇAMENTOS.', ncols=10)
    sh.widths({**{CL(c): 14.5 for c in range(2, 12)}})
    ws.sheet_properties.tabColor = COR['header']
    sh.put('B6', 'Mês:', 'lblb', h='right')
    sh.put('C6', None, 'in', bold=True, h='center')
    sh.put('D6', 'Ano:', 'lblb', h='right')
    sh.put('E6', None, 'in', INT, bold=True, h='center')
    sh.dv('C6', '=Lista_Meses')
    sh.dv('E6', '=Lista_Anos')
    ctx.name('Dash_Mes', S, 'C6')
    ctx.name('Dash_Ano', S, 'E6')
    ctx.name('Exemplo_DashSel', S, 'C6:E6')
    sh.put('F6', '⬅ digite aqui (vazio = mês atual de CONFIGURAÇÕES)', 'note', color=COR['azul'], bold=True)
    sh.put('B7', '="Olá"&IF(Cfg_Nome<>"",", "&Cfg_Nome,"")&"! Período: "&Sel_Label&IF(COUNTIF(L_Obs,"EXEMPLO*")>0,"   ·   ⚠ DADOS DE EXEMPLO ATIVOS (veja COMEÇAR AQUI para limpar)","")',
           'plain', bold=True, color=COR['header'])
    sh.merge('B7:K7')
    sh.cf_formula('B7', 'ISNUMBER(SEARCH("EXEMPLO",B7))', COR['amarelo_l'], '9A5B00', True)
    sh.put('B9', 'Indicadores do mês', 'sec')

    def card(r, c, label, formula, fmt, color, sub, cf=None):
        L1, L2 = CL(c), CL(c + 1)
        for rr in (r, r + 1, r + 2):
            sh.merge(f'{L1}{rr}:{L2}{rr}')
            for L in (L1, L2):
                sh.put(f'{L}{rr}', None, 'calc', fillc=COR['cinza'], border=False)
        sh.put(f'{L1}{r}', label, 'calc', bold=True, size=10, color=COR['mudo'], h='center', border=False)
        sh.put(f'{L1}{r + 1}', formula, 'calc', fmt, bold=True, size=18, color=color, h='center', border=False)
        sh.put(f'{L1}{r + 2}', sub, 'calc', size=9, color=COR['mudo'], h='center', border=False)
        ws.row_dimensions[r + 1].height = 30
        # moldura
        for rr in (r, r + 1, r + 2):
            for L in (L1, L2):
                cell = ws[f'{L}{rr}']
                cell.border = Border(left=thin if L == L1 else None, right=thin if L == L2 else None,
                                     top=Side(style='medium', color=color) if rr == r else None,
                                     bottom=thin if rr == r + 2 else None)
        if cf:
            for fx, col in cf:
                sh.cf_formula(f'{L1}{r + 1}', fx.replace('#', f'{L1}{r + 1}'), None, col, True)

    def vs(cur, prev):
        return f'IF({prev}=0,"sem mês anterior",IF({cur}>={prev},"▲ ","▼ ")&FIXED(ABS(({cur}-{prev})/ABS({prev}))*100,1)&"% vs mês anterior")'

    card(10, 2, 'RECEITAS', '=K_Rec', BRLZ, COR['verde'], f'={vs("K_Rec", "K_RecAnt")}')
    card(10, 4, 'DESPESAS', '=K_Desp', BRLZ, COR['vermelho'], f'={vs("K_Desp", "K_DespAnt")}')
    card(10, 6, 'SALDO DO MÊS', '=K_Saldo', BRLZ, COR['verde'], '="Receitas − despesas"',
         [('#<0', COR['vermelho'])])
    card(10, 8, 'TAXA DE POUPANÇA', '=K_Poup', PCT, COR['azul'], '="Meta: "&ROUND(K_PoupMeta*100,0)&"%"',
         [('#<K_PoupMeta', COR['amarelo']), ('#<0', COR['vermelho'])])
    card(10, 10, 'GASTO MÉDIO DIÁRIO', '=K_GastoDia', BRLZ, COR['vermelho'], '="Despesas ÷ dias do período"')
    card(14, 2, 'PATRIMÔNIO LÍQUIDO', '=Pat_Liquido', BRLZ, COR['azul'], '="Ativos R$ "&FIXED(Pat_Ativos,0)&" · Passivos R$ "&FIXED(Pat_Passivos,0)',
         [('#<0', COR['vermelho'])])
    card(14, 4, 'TOTAL DE DÍVIDAS', '=Div_Total', BRLZ, COR['vermelho'], '="Parcelas = "&ROUND(Div_PctParc*100,1)&"% da renda"')
    card(14, 6, 'RESERVA DE EMERGÊNCIA', '=K_ReservaPct', PCT, COR['roxo'],
         '=FIXED(K_ReservaMeses,1)&" de "&K_MesesAlvo&" meses cobertos"', [('#<0.5', COR['vermelho']), ('#<1', COR['amarelo'])])
    card(14, 8, 'DIAS DE CAIXA RESTANTES', '=K_DiasCaixa', INT, COR['azul'], '="Saldo em contas R$ "&FIXED(K_SaldoContas,0)',
         [('#<30', COR['vermelho']), ('#<60', COR['amarelo'])])
    card(14, 10, 'META DO MÊS (APORTE)', '=K_AportePct', PCT, COR['roxo'],
         '="Guardado R$ "&FIXED(K_Apo,0)&" de R$ "&FIXED(K_AporteMeta,0)', [('#<0.7', COR['vermelho']), ('#<1', COR['amarelo']), ('#>=1', COR['verde'])])
    # ---- saúde financeira
    sh.put('B18', 'Saúde financeira', 'sec')
    sh.merge('B19:C22')
    sh.put('B19', '=IF(K_Rec+K_Desp=0,"—",H_Nota)', 'calc', INT, bold=True, size=36, h='center', border=True)
    for rr in range(19, 23):
        for L in 'BC':
            ws[f'{L}{rr}'].border = BORDER
            ws[f'{L}{rr}'].fill = fill(COR['cinza'])
    sh.merge('D19:E19')
    sh.put('D19', '="Nota de 0 a 100"', 'calc', bold=True, h='center', color=COR['mudo'])
    sh.merge('D20:E22')
    sh.put('D20', '=IF(NOT(ISNUMBER(B19)),"Aguardando lançamentos",IF(H_Nota>=70,"🟢 SAUDÁVEL",IF(H_Nota>=40,"🟡 ATENÇÃO","🔴 CRÍTICO")))', 'calc', bold=True, size=14, h='center', wrap=True)
    for rr in range(19, 23):
        for L in 'DE':
            ws[f'{L}{rr}'].border = BORDER
            ws[f'{L}{rr}'].fill = fill(COR['cinza'])
    for fx, bg, fg in (('AND(ISNUMBER($B$19),$B$19>=70)', COR['verde_l'], COR['verde']),
                       ('AND(ISNUMBER($B$19),$B$19>=40,$B$19<70)', COR['amarelo_l'], '9A5B00'),
                       ('AND(ISNUMBER($B$19),$B$19<40)', COR['vermelho_l'], COR['vermelho'])):
        sh.cf_formula('B19:E22', fx, bg, fg, True)
    for i, h in enumerate(['Componente', 'Pontos', 'Máx.', 'Situação']):
        L = ['F', 'I', 'J', 'K'][i]
        sh.put(f'{L}18', h, 'hdr')
    sh.merge('F18:H18')
    comp = [('Taxa de poupança', 'H_Poup'), ('Endividamento sobre a renda', 'H_Div'), ('Reserva de emergência', 'H_Res'),
            ('Gastos dentro do orçamento', 'H_Orc'), ('Regularidade de aportes', 'H_Apo')]
    for i, (lb, nm) in enumerate(comp):
        r = 19 + i
        sh.merge(f'F{r}:H{r}')
        sh.put(f'F{r}', lb, 'calc')
        sh.put(f'I{r}', f'={nm}', 'calc', NUM1, h='center', bold=True)
        sh.put(f'J{r}', 20, 'calc', INT, h='center')
        sh.put(f'K{r}', f'=IF(I{r}>=14,"🟢",IF(I{r}>=8,"🟡","🔴"))', 'calc', h='center')
        sh.cf_formula(f'I{r}', f'I{r}>=14', COR['verde_l'], COR['verde'], True)
        sh.cf_formula(f'I{r}', f'AND(I{r}>=8,I{r}<14)', COR['amarelo_l'], '9A5B00', True)
        sh.cf_formula(f'I{r}', f'I{r}<8', COR['vermelho_l'], COR['vermelho'], True)
    # ---- alertas
    sh.put('B25', 'Alertas automáticos', 'sec')
    for k in range(1, 9):
        r = 25 + k
        sh.merge(f'B{r}:K{r}')
        sh.put(f'B{r}', f'=IF({k}=1,IF(Alert_Total=0,"✔ Nenhum alerta no momento. Tudo em ordem!",IFERROR(INDEX(Alert_Msg,MATCH(1,Alert_Cnt,0)),"")),'
               f'IFERROR(INDEX(Alert_Msg,MATCH({k},Alert_Cnt,0)),""))', 'calc', h='left', border=False)
        sh.cf_formula(f'B{r}:K{r}', f'LEFT($B{r},2)="🔴"', COR['vermelho_l'], COR['vermelho'], True)
        sh.cf_formula(f'B{r}:K{r}', f'LEFT($B{r},2)="🟡"', COR['amarelo_l'], '9A5B00', True)
        sh.cf_formula(f'B{r}:K{r}', f'LEFT($B{r},1)="✔"', COR['verde_l'], COR['verde'], True)
        sh.cf_formula(f'B{r}:K{r}', f'$B{r}=""', COR['cinza'], COR['cinza'])
    sh.put('B34', '="Mostrando "&MIN(8,Alert_Total)&" de "&Alert_Total&" alerta(s)."', 'note')
    # ---- comparativos
    sh.put('B36', 'Comparativos', 'sec')
    for L, t in zip('BDEFGHI', ['Indicador', 'Mês atual', 'Mês anterior', 'Variação', 'Mesmo mês, ano anterior', 'Variação']):
        pass
    heads = [('B', 'Indicador'), ('D', 'Mês atual'), ('E', 'Mês anterior'), ('F', 'Variação'), ('G', 'Ano anterior'), ('H', 'Variação')]
    for L, t in heads:
        sh.put(f'{L}37', t, 'hdr')
    sh.merge('B37:C37')
    sh.put('C37', None, 'hdr')
    comp2 = [('Receitas', 'K_Rec', 'K_RecAnt', 'K_RecAA', BRLZ, 'up'), ('Despesas', 'K_Desp', 'K_DespAnt', 'K_DespAA', BRLZ, 'down'),
             ('Saldo do mês', 'K_Saldo', 'K_SaldoAnt', 'K_SaldoAA', BRLZ, 'up'), ('Aportes', 'K_Apo', 'K_ApoAnt', 'K_ApoAA', BRLZ, 'up'),
             ('Taxa de poupança', 'K_Poup', 'K_PoupAnt', 'K_PoupAA', PCT, 'up')]
    for i, (lb, cur, ant, aa, fm, sense) in enumerate(comp2):
        r = 38 + i
        sh.merge(f'B{r}:C{r}')
        sh.put(f'B{r}', lb, 'calc', bold=True)
        sh.put(f'C{r}', None, 'calc')
        sh.put(f'D{r}', f'={cur}', 'calc', fm)
        sh.put(f'E{r}', f'={ant}', 'calc', fm)
        sh.put(f'G{r}', f'={aa}', 'calc', fm)
        for L, base in (('F', f'E{r}'), ('H', f'G{r}')):
            if fm == PCT:
                f = f'=IF(AND({base}=0,D{r}=0),"—",IF(D{r}>={base},"▲ ","▼ ")&FIXED(ABS(D{r}-{base})*100,1)&" p.p.")'
            else:
                f = f'=IF({base}=0,"—",IF(D{r}>={base},"▲ ","▼ ")&FIXED(ABS((D{r}-{base})/ABS({base}))*100,1)&"%")'
            sh.put(f'{L}{r}', f, 'calc', h='center', bold=True)
            good = COR['verde'] if sense == 'up' else COR['vermelho']
            bad = COR['vermelho'] if sense == 'up' else COR['verde']
            sh.cf_formula(f'{L}{r}', f'LEFT({L}{r},1)="▲"', None, good, True)
            sh.cf_formula(f'{L}{r}', f'LEFT({L}{r},1)="▼"', None, bad, True)
    # ---- minhas metas
    sh.put('B44', 'Minhas metas (5 principais em andamento)', 'sec')
    for L, t in (('B', 'Meta'), ('E', 'Progresso'), ('H', '% concluído'), ('I', 'Status'), ('J', 'Falta')):
        sh.put(f'{L}45', t, 'hdr')
    sh.merge('B45:D45')
    sh.merge('E45:G45')
    sh.merge('J45:K45')
    for k in range(1, 6):
        r = 45 + k
        ix = f'IFERROR(MATCH({k},Prio_Ordem,0),0)'
        sh.merge(f'B{r}:D{r}')
        sh.merge(f'E{r}:G{r}')
        sh.merge(f'J{r}:K{r}')
        sh.put(f'L{r}', f'={ix}', 'calc', INT, size=8, color=COR['cinza'], border=False)
        sh.put(f'B{r}', f'=IF($L{r}=0,"",INDEX(Meta_Nome,$L{r}))', 'calc', bold=True)
        sh.put(f'E{r}', f'=IF($L{r}=0,"",INDEX(Meta_Barra,$L{r}))', 'calc', size=10)
        sh.put(f'H{r}', f'=IF($L{r}=0,"",INDEX(Meta_Pct,$L{r}))', 'calc', PCT, h='center')
        sh.put(f'I{r}', f'=IF($L{r}=0,"",INDEX(Meta_Status,$L{r}))', 'calc', h='center', bold=True)
        sh.put(f'J{r}', f'=IF($L{r}=0,"",INDEX(Meta_Falta,$L{r}))', 'calc', BRL)
        for st, cr, fg in (('Concluída', COR['verde'], 'FFFFFF'), ('No prazo', COR['azul'], 'FFFFFF'),
                           ('Atenção', COR['amarelo'], '1F2933'), ('Atrasada', COR['vermelho'], 'FFFFFF')):
            sh.cf_formula(f'I{r}', f'$I{r}="{st}"', cr, fg, True)
            sh.cf_formula(f'E{r}', f'$I{r}="{st}"', None, cr)
    sh.put('B51', 'Veja detalhes, ritmo e planejador em METAS E OBJETIVOS.', 'note')
    # ---- gráficos
    sh.put('B53', 'Gráficos', 'sec')
    build_charts(ctx, ws, 54)
    sh.finish(freeze='A8')
    ws.print_area = 'A1:L125'
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_setup.scale = 68
    ws.row_breaks.append(Break(id=43))
    ws.row_breaks.append(Break(id=87))
    return sh


# ---------------------------------------------------------------- COMEÇAR AQUI
def build_inicio(ctx):
    sh = Sh(ctx, 'COMEÇAR AQUI')
    ws = sh.ws
    sh.header('COMEÇAR AQUI', 'Bem-vindo(a)! Em 2 minutos você entende onde digitar e onde olhar. Digite em LANÇAMENTOS; o resto é automático.',
              ncols=9)
    sh.widths({'B': 6, 'C': 46, 'D': 44, 'E': 16, 'F': 14, 'G': 14, 'H': 14, 'I': 14})
    sh.put('B5', '=IF(COUNTIF(L_Obs,"EXEMPLO*")>0,"⚠ Esta planilha contém DADOS DE EXEMPLO fictícios (marcados \'EXEMPLO\'). Veja no fim desta aba como limpar tudo.","✔ Sem dados de exemplo: planilha pronta para os seus dados.")',
           'plain', bold=True)
    sh.merge('B5:I5')
    sh.cf_formula('B5', 'LEFT(B5,1)="⚠"', COR['amarelo_l'], '9A5B00', True)
    sh.cf_formula('B5', 'LEFT(B5,1)="✔"', COR['verde_l'], COR['verde'], True)
    sh.put('B7', 'Passo a passo', 'sec')
    for L, t in (('B', '#'), ('C', 'O que fazer'), ('D', 'Detalhe'), ('E', 'Ir para'), ('F', 'Feito?')):
        sh.put(f'{L}8', t, 'hdr')
    passos = [
        ('1', 'Configurar', 'Em CONFIGURAÇÕES: nome, renda mensal líquida, meta de poupança e metas de comportamento.',
         'CONFIGURAÇÕES', '=IF(AND(Cfg_Nome<>"",N(Cfg_Renda)>0),"✔","○")'),
        ('2', 'Cadastrar contas e cartões', 'Também em CONFIGURAÇÕES: contas (com saldo inicial), cartões (limite, fechamento, vencimento), categorias.',
         'CONFIGURAÇÕES', '=IF(COUNTA(Contas_Nome)>0,"✔","○")'),
        ('3', 'Lançar no dia a dia', 'Só em LANÇAMENTOS: uma linha por receita, despesa, transferência ou aporte. Parcelado = 1 linha.',
         'LANÇAMENTOS', '=IF(COUNT(L_Data)>0,"✔","○")'),
        ('4', 'Conferir o Dashboard', 'Escolha o mês no topo do DASHBOARD e veja indicadores, nota de saúde, alertas e gráficos.',
         'DASHBOARD', '=IF(K_Rec+K_Desp>0,"✔","○")'),
        ('5', 'Revisar todo mês', 'Atualize ORÇAMENTO, METAS, retrato do PATRIMÔNIO e confira o CALENDÁRIO DE CONTAS.',
         'ORÇAMENTO', '=IF(AND(SUM(Orc_TotDesp)>0,COUNTA(Meta_Nome)>0),"✔","○")'),
    ]
    for i, (n, t, d, tgt, chk) in enumerate(passos):
        r = 9 + i
        sh.put(f'B{r}', n, 'calc', bold=True, h='center', size=14, color=COR['azul'])
        sh.put(f'C{r}', t, 'calc', bold=True)
        sh.put(f'D{r}', d, 'calc', wrap=True)
        c = sh.put(f'E{r}', f'➜ {tgt.title() if False else tgt}', 'calc', h='center', color=COR['azul'], size=10)
        c.hyperlink = Hyperlink(ref=f'E{r}', location=f"'{tgt}'!A1", display=tgt)
        c.font = fnt(10, color=COR['azul'], underline='single')
        sh.put(f'F{r}', chk, 'calc', h='center', bold=True, size=14)
        ws.row_dimensions[r].height = 36
    sh.cf_formula('F9:F13', 'F9="✔"', COR['verde_l'], COR['verde'], True)
    sh.cf_formula('F9:F13', 'F9="○"', COR['amarelo_l'], COR['amarelo'], True)
    # checklist
    sh.put('B16', 'Checklist de configuração inicial (✔ automático)', 'sec')
    for L, t in (('B', ''), ('C', 'Item'), ('D', 'Como saber que está pronto'), ('E', 'Ir para'), ('F', 'Status')):
        sh.put(f'{L}17', t, 'hdr')
    itens = [
        ('Informar seu nome', 'CONFIGURAÇÕES: célula "Nome"', 'CONFIGURAÇÕES', '=IF(Cfg_Nome<>"","✔","○")'),
        ('Informar a renda mensal líquida', 'CONFIGURAÇÕES: "Renda mensal líquida média"', 'CONFIGURAÇÕES', '=IF(N(Cfg_Renda)>0,"✔","○")'),
        ('Definir meses da reserva e % de poupança', 'CONFIGURAÇÕES: padrões 6 meses e 20%', 'CONFIGURAÇÕES', '=IF(AND(N(Cfg_MesesReserva)>0,N(Cfg_MetaPoup)>0),"✔","○")'),
        ('Cadastrar ao menos 1 conta bancária', 'CONFIGURAÇÕES, seção 3', 'CONFIGURAÇÕES', '=IF(COUNTA(Contas_Nome)>0,"✔","○")'),
        ('Cadastrar cartões (se usar)', 'CONFIGURAÇÕES, seção 4', 'CONFIGURAÇÕES', '=IF(COUNTA(Cartoes_Nome)>0,"✔","○")'),
        ('Revisar categorias de despesa e receita', 'CONFIGURAÇÕES, seções 5 e 6', 'CONFIGURAÇÕES', '=IF(COUNTA(CatD_Nome)>0,"✔","○")'),
        ('Definir o orçamento mensal', 'ORÇAMENTO, seção 1', 'ORÇAMENTO', '=IF(SUM(Orc_TotDesp)>0,"✔","○")'),
        ('Cadastrar recorrentes (salário, aluguel, assinaturas)', 'RECORRENTES', 'RECORRENTES', '=IF(COUNTA(Rec_Desc)>0,"✔","○")'),
        ('Cadastrar ao menos 1 meta', 'METAS E OBJETIVOS', 'METAS E OBJETIVOS', '=IF(COUNTA(Meta_Nome)>0,"✔","○")'),
        ('Cadastrar dívidas (se houver)', 'DÍVIDAS E FINANCIAMENTOS', 'DÍVIDAS E FINANCIAMENTOS', '=IF(COUNTA(Div_Nome)>0,"✔","○")'),
        ('Cadastrar investimentos (se houver)', 'INVESTIMENTOS', 'INVESTIMENTOS', '=IF(COUNTA(Inv_Nome)>0,"✔","○")'),
        ('Fazer o primeiro lançamento', 'LANÇAMENTOS', 'LANÇAMENTOS', '=IF(COUNT(L_Data)>0,"✔","○")'),
    ]
    for i, (t, d, tgt, chk) in enumerate(itens):
        r = 18 + i
        sh.put(f'B{r}', i + 1, 'calc', h='center')
        sh.put(f'C{r}', t, 'calc')
        sh.put(f'D{r}', d, 'calc')
        c = sh.put(f'E{r}', f'➜ {tgt}', 'calc', h='center', size=10)
        c.hyperlink = Hyperlink(ref=f'E{r}', location=f"'{tgt}'!A1", display=tgt)
        c.font = fnt(10, color=COR['azul'], underline='single')
        sh.put(f'F{r}', chk, 'calc', h='center', bold=True, size=13)
    last = 18 + len(itens) - 1
    sh.cf_formula(f'F18:F{last}', 'F18="✔"', COR['verde_l'], COR['verde'], True)
    sh.cf_formula(f'F18:F{last}', 'F18="○"', COR['amarelo_l'], COR['amarelo'], True)
    sh.put(f'C{last + 1}', 'Progresso da configuração', 'calc', bold=True)
    sh.put(f'D{last + 1}', f'=COUNTIF(F18:F{last},"✔")&" de {len(itens)} itens ("&ROUND(COUNTIF(F18:F{last},"✔")/{len(itens)}*100,0)&"%)"', 'calc', bold=True)
    # legenda de cores
    z = last + 4
    sh.put(f'B{z}', 'Legenda de cores', 'sec')
    leg = [(COR['azulclaro'], COR['texto'], 'Azul claro = célula de ENTRADA (você digita)'),
           (COR['cinza'], COR['texto'], 'Cinza claro = célula CALCULADA e protegida (não digite)'),
           (COR['verde'], 'FFFFFF', 'Verde = receitas, saldo positivo, dentro do orçamento, meta atingida'),
           (COR['vermelho'], 'FFFFFF', 'Vermelho = despesas, saldo negativo, estouro, dívidas'),
           (COR['amarelo'], COR['texto'], 'Amarelo/laranja = atenção, perto do limite, pendente'),
           (COR['azul'], 'FFFFFF', 'Azul = investimentos, patrimônio, neutro, títulos'),
           (COR['roxo'], 'FFFFFF', 'Roxo = metas e planejamento')]
    for i, (bg, fg, tx) in enumerate(leg):
        r = z + 1 + i
        sh.put(f'B{r}', None, 'calc', fillc=bg)
        sh.put(f'C{r}', tx, 'lbl')
    # mapa das abas
    m = z + len(leg) + 3
    sh.put(f'B{m}', 'Mapa das abas (clique para ir)', 'sec')
    mapa = [('DASHBOARD', 'Visão geral: indicadores, nota de saúde, alertas, comparativos e gráficos'),
            ('LANÇAMENTOS', 'ÚNICA aba de digitação diária'), ('RECORRENTES', 'Itens que se repetem (alimenta as projeções)'),
            ('ORÇAMENTO', 'Planejado x realizado por categoria e regra 50/30/20'), ('FLUXO DE CAIXA', 'Matriz de 12 meses, médias e sazonalidade'),
            ('PROJEÇÕES E CENÁRIOS', '12/24 meses, 3 cenários e simulador "E se?"'), ('CARTÕES DE CRÉDITO', 'Limite, faturas e parcelas futuras'),
            ('DÍVIDAS E FINANCIAMENTOS', 'Amortização Price/SAC e simulador Bola de Neve x Avalanche'),
            ('INVESTIMENTOS', 'Carteira, rentabilidade e rebalanceamento'), ('METAS E OBJETIVOS', 'Metas, planejador reverso e prioridades'),
            ('PATRIMÔNIO LÍQUIDO', 'Ativos − passivos e histórico mensal'), ('CALENDÁRIO DE CONTAS', 'Vencimentos e total a pagar em 7/15/30 dias'),
            ('IMPOSTO DE RENDA', 'Organização para a declaração'), ('HISTÓRICO E ANÁLISES', 'Rankings, evolução, tags e assinaturas'),
            ('BACKUP', 'Como guardar e versionar o arquivo; premissas e limitações'), ('CONFIGURAÇÕES', 'Parâmetros e cadastros')]
    for i, (aba, d) in enumerate(mapa):
        r = m + 1 + i
        c = sh.put(f'C{r}', aba, 'calc', bold=True)
        c.hyperlink = Hyperlink(ref=f'C{r}', location=f"'{aba}'!A1", display=aba)
        c.font = fnt(11, bold=True, color=COR['azul'], underline='single')
        sh.put(f'D{r}', d, 'calc')
    # limpar exemplos
    cl = m + len(mapa) + 3
    sh.put(f'B{cl}', 'Como limpar os dados de EXEMPLO', 'sec')
    sh.put(f'B{cl + 1}', 'Opção 1 (um passo): use o arquivo "Planilha_Financeira_Avancada_LIMPA.xlsx", entregue junto com este — mesma planilha, sem exemplos.', 'lblb')
    sh.put(f'B{cl + 2}', 'Opção 2 (neste arquivo): em cada aba abaixo, digite o NOME no campo "Caixa de Nomes" (à esquerda da barra de fórmulas), tecle Enter e depois Delete.', 'lbl')
    for i, h in enumerate(['Aba', 'Nome para digitar na Caixa de Nomes', 'O que apaga']):
        sh.put(f'{CL(3 + i)}{cl + 3}', h, 'hdr')
    lim = [('LANÇAMENTOS', 'Exemplo_Lancamentos', 'Todas as linhas azuis de lançamentos'), ('RECORRENTES', 'Exemplo_Recorrentes', 'Recorrentes'),
           ('ORÇAMENTO', 'Exemplo_OrcDespesas  e  Exemplo_OrcReceitas', 'Valores planejados'), ('METAS E OBJETIVOS', 'Exemplo_Metas  e  Exemplo_Planejador', 'Metas e planejador'),
           ('DÍVIDAS E FINANCIAMENTOS', 'Exemplo_Dividas  e  Exemplo_SimExtra', 'Dívidas e extra do simulador'),
           ('INVESTIMENTOS', 'Exemplo_Invest  e  Exemplo_InvAlvo', 'Ativos e alvos'), ('PATRIMÔNIO LÍQUIDO', 'Exemplo_PatAtivos, Exemplo_PatPassivo, Exemplo_PatAno0, Exemplo_PatHist', 'Bens e histórico'),
           ('IMPOSTO DE RENDA', 'Exemplo_IR_Trib, Exemplo_IR_Isent, Exemplo_IR_Excl, Exemplo_IR_Bens, Exemplo_IR_Div', 'Dados de IR'),
           ('PROJEÇÕES E CENÁRIOS', 'Exemplo_Simulador', 'Campos do simulador'),
           ('CONFIGURAÇÕES', 'Apague manualmente nome, renda, contas, cartões e limites das categorias', 'Cadastros pessoais'),
           ('DASHBOARD', 'Exemplo_DashSel', 'Mês/ano escolhidos (volta ao mês atual)')]
    for i, (a, n, w) in enumerate(lim):
        r = cl + 4 + i
        sh.put(f'C{r}', a, 'calc')
        sh.put(f'D{r}', n, 'calc', wrap=True)
        sh.put(f'E{r}', w, 'calc', wrap=True)
        ws.row_dimensions[r].height = 30
    sh.merge(f'B{cl + 1}:I{cl + 1}')
    sh.merge(f'B{cl + 2}:I{cl + 2}')
    # dicas
    d0 = cl + len(lim) + 6
    sh.put(f'B{d0}', 'Regras de ouro (leia uma vez)', 'sec')
    dicas = [
        'Valor em LANÇAMENTOS = valor TOTAL da compra. Parcelado: informe "Parcela atual" e "Total de parcelas" (ex.: 1 e 10) — as demais parcelas surgem sozinhas nos meses futuros.',
        'Compra no cartão: escolha o cartão na coluna "Conta ou Cartão". A fatura de referência é calculada pelo dia de fechamento.',
        'Pagou a fatura? Lance como Transferência (valor NEGATIVO) na conta de onde saiu o dinheiro — assim o saldo da conta fica correto sem contar a despesa duas vezes.',
        'Transferência entre contas: lance duas linhas — uma negativa na conta de origem e uma positiva na de destino.',
        'Aporte em investimento/meta: Tipo = Investimento; se for para uma meta, escolha em "Meta vinculada". O valor guardado da meta soma sozinho.',
        'Não cadastre o mesmo item em RECORRENTES e em lançamentos futuros (as projeções somariam duas vezes).',
        'Listas (contas, categorias, tags...) devem ser preenchidas de cima para baixo, sem linhas em branco no meio.',
        'As abas calculadas estão protegidas SEM senha (Revisão > Desproteger Planilha) — só as células azuis aceitam digitação.',
    ]
    for i, t in enumerate(dicas):
        r = d0 + 1 + i
        sh.put(f'B{r}', i + 1, 'calc', h='center')
        sh.put(f'C{r}', t, 'calc', wrap=True)
        sh.merge(f'C{r}:I{r}')
        ws.row_dimensions[r].height = 32
    sh.finish(freeze='A5', landscape=False)
    return sh


def build_backup(ctx):
    sh = Sh(ctx, 'BACKUP')
    ws = sh.ws
    sh.header('BACKUP E VERSIONAMENTO', 'Seus dados vivem neste arquivo: proteja-o. Instruções de backup, controle de versões, premissas e limitações conhecidas.',
              ncols=6)
    sh.widths({'B': 4, 'C': 90, 'D': 18, 'E': 18})
    sh.put('B5', 'Controle de backup', 'sec')
    sh.put('C6', 'Data do último backup (digite aqui):', 'lblb', h='right')
    sh.put('D6', None, 'in', DATA)
    sh.put('C7', 'Dias desde o último backup:', 'lbl', h='right')
    sh.put('D7', '=IF(ISNUMBER(D6),Cfg_Hoje-D6,"—")', 'calc', INT, h='center')
    sh.put('E7', '=IF(NOT(ISNUMBER(D7)),"⚠ faça um backup",IF(D7>30,"⚠ backup atrasado","✔ em dia"))', 'calc', bold=True)
    sh.cf_formula('E7', 'LEFT(E7,1)="⚠"', COR['amarelo_l'], '9A5B00', True)
    sh.cf_formula('E7', 'LEFT(E7,1)="✔"', COR['verde_l'], COR['verde'], True)
    ctx.name('Exemplo_Backup', 'BACKUP', 'D6')
    txt = [
        ('Como fazer backup', None),
        ('1. Uma vez por mês (de preferência no dia em que você faz o fechamento), use Arquivo > Salvar como e grave uma CÓPIA com a data no nome. Ex.: Planilha_Financeira_2026-10.xlsx.', None),
        ('2. Guarde a cópia em um lugar diferente do original: nuvem (Google Drive, OneDrive, iCloud, Dropbox) e, de vez em quando, um pen drive ou HD externo.', None),
        ('3. Ative o histórico de versões da nuvem (Excel/OneDrive: Arquivo > Informações > Histórico de versões). Ele permite voltar a um estado anterior se você apagar algo sem querer.', None),
        ('4. Antes de mudanças grandes (limpar exemplos, colar muitas linhas, trocar de ano), faça uma cópia extra.', None),
        ('5. Uma vez por ano, exporte LANÇAMENTOS para CSV (Salvar como > CSV) e arquive junto com os comprovantes do Imposto de Renda.', None),
        ('Versionamento sugerido', None),
        ('• Nome do arquivo de trabalho: Planilha_Financeira_Avancada.xlsx (sempre o mesmo). Cópias mensais: _AAAA-MM. Ao virar o ano, salve uma cópia "_AAAA_fechado" e continue no mesmo arquivo (a planilha trata vários anos).', None),
        ('• Anote em um caderno ou neste campo quando fizer mudanças estruturais (novas categorias, novos cartões).', None),
        ('Reexibir abas ocultas', None),
        ('• As abas TABELAS, CÁLCULOS e SIM_DÍVIDAS ficam ocultas (apoio). Para ver: clique com o botão direito em qualquer aba > Reexibir.', None),
        ('Proteção', None),
        ('• As abas estão protegidas SEM senha para evitar apagar fórmulas sem querer. Para editar uma célula calculada: Revisão > Desproteger Planilha. Lembre de proteger de novo depois.', None),
        ('Premissas adotadas', None),
        ('• Realizado = todos os lançamentos do mês (qualquer status, regime de competência pela data). O saldo das contas considera só os lançamentos PAGOS.', None),
        ('• Orçamento planejado vale para qualquer ano (12 colunas de meses). Ajuste os valores ao virar o ano, se mudarem.', None),
        ('• Compras no cartão entram nos gastos pela data da compra (parcelas mês a mês); a fatura de referência segue o dia de fechamento.', None),
        ('• Reserva de emergência = soma do valor guardado das metas da categoria "Reserva". Poupança = (receitas − despesas) ÷ receitas.', None),
        ('• Projeções: receitas e despesas fixas vêm de RECORRENTES (se não houver, usa renda configurada e média histórica); despesas variáveis = média dos últimos 3 meses fechados corrigida pela inflação.', None),
        ('• Ritmo de aporte das metas = média dos 3 últimos meses fechados anteriores ao mês atual de CONFIGURAÇÕES.', None),
        ('Limitações conhecidas', None),
        ('• Capacidade: 1.500 lançamentos, 40 recorrentes, 15 metas, 8 cartões, 8 dívidas, 30 ativos, 18 categorias de despesa e 9 de receita (aumentar exige estender os intervalos).', None),
        ('• Sem macros: "copiar orçamento do mês anterior" e "limpar exemplos" são feitos com Ctrl+C/Ctrl+V e a Caixa de Nomes.', None),
        ('• Gráficos usam as 8-10 maiores categorias; categorias novas aparecem automaticamente se estiverem entre as maiores do mês.', None),
        ('• Simulador de dívidas: horizonte de 10 anos, ordem de ataque fixa (calculada pelo saldo/juros iniciais).', None),
        ('• Imposto de Renda: apenas organização; limites de dedução são valores de referência editáveis.', None),
        ('• Funciona no Excel, LibreOffice e Google Planilhas; algumas formatações condicionais e barras de dados podem variar entre programas.', None),
    ]
    r = 9
    for t, _ in txt:
        head = not t.startswith(('•', '1.', '2.', '3.', '4.', '5.'))
        sh.put(f'C{r}', t, 'sec' if head else 'lbl', wrap=True)
        if not head:
            ws.row_dimensions[r].height = 32
        r += 1
    sh.finish(freeze='A5', landscape=False)
    return sh
