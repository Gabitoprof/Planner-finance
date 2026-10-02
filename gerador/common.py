"""Utilitários comuns do gerador da Planilha Financeira Avançada."""
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.formatting.rule import FormulaRule, CellIsRule, DataBarRule, ColorScaleRule

COR = dict(verde='2E9E5B', vermelho='D64545', amarelo='F2A93B', azul='2F6FD0', roxo='7B5EA7',
           cinza='F2F4F7', azulclaro='E8F0FE', texto='1F2933', header='1B3A6B', borda='D0D5DD',
           verde_l='DDF3E4', vermelho_l='FADBD8', amarelo_l='FDEFD2', azul_l='DCE8FA', roxo_l='E9E1F3',
           mudo='667085', branco='FFFFFF')

BRL = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
BRL0 = '"R$" #,##0;-"R$" #,##0;"-"'
BRLZ = '"R$" #,##0.00;-"R$" #,##0.00'
PCT = '0.0%'
DATA = 'dd/mm/yyyy'
MESANO = 'mm/yyyy'
INT = '0'
NUM1 = '0.0'

NL = 1500          # linhas de lançamentos
LR0 = 7            # primeira linha de dados em LANÇAMENTOS
LR1 = LR0 + NL - 1

SHEETS = ['COMEÇAR AQUI', 'CONFIGURAÇÕES', 'LANÇAMENTOS', 'RECORRENTES', 'DASHBOARD', 'ORÇAMENTO',
          'FLUXO DE CAIXA', 'PROJEÇÕES E CENÁRIOS', 'CARTÕES DE CRÉDITO', 'DÍVIDAS E FINANCIAMENTOS',
          'INVESTIMENTOS', 'METAS E OBJETIVOS', 'PATRIMÔNIO LÍQUIDO', 'CALENDÁRIO DE CONTAS',
          'IMPOSTO DE RENDA', 'HISTÓRICO E ANÁLISES', 'TABELAS', 'CÁLCULOS', 'SIM_DÍVIDAS', 'BACKUP']

NAV = [('◀ Voltar ao Dashboard', 'DASHBOARD'), ('Começar aqui', 'COMEÇAR AQUI'),
       ('Configurações', 'CONFIGURAÇÕES'), ('Lançamentos', 'LANÇAMENTOS'), ('Recorrentes', 'RECORRENTES'),
       ('Orçamento', 'ORÇAMENTO'), ('Fluxo de caixa', 'FLUXO DE CAIXA'), ('Projeções', 'PROJEÇÕES E CENÁRIOS'),
       ('Cartões', 'CARTÕES DE CRÉDITO'), ('Dívidas', 'DÍVIDAS E FINANCIAMENTOS'),
       ('Investimentos', 'INVESTIMENTOS'), ('Metas', 'METAS E OBJETIVOS'), ('Patrimônio', 'PATRIMÔNIO LÍQUIDO'),
       ('Calendário', 'CALENDÁRIO DE CONTAS'), ('Imposto de Renda', 'IMPOSTO DE RENDA'),
       ('Análises', 'HISTÓRICO E ANÁLISES'), ('Backup', 'BACKUP')]

TAB_COLORS = {'COMEÇAR AQUI': COR['verde'], 'CONFIGURAÇÕES': COR['mudo'], 'LANÇAMENTOS': COR['azul'],
              'RECORRENTES': COR['azul'], 'DASHBOARD': COR['header'], 'ORÇAMENTO': COR['amarelo'],
              'FLUXO DE CAIXA': COR['verde'], 'PROJEÇÕES E CENÁRIOS': COR['roxo'],
              'CARTÕES DE CRÉDITO': COR['vermelho'], 'DÍVIDAS E FINANCIAMENTOS': COR['vermelho'],
              'INVESTIMENTOS': COR['azul'], 'METAS E OBJETIVOS': COR['roxo'], 'PATRIMÔNIO LÍQUIDO': COR['azul'],
              'CALENDÁRIO DE CONTAS': COR['amarelo'], 'IMPOSTO DE RENDA': COR['mudo'],
              'HISTÓRICO E ANÁLISES': COR['azul'], 'TABELAS': COR['mudo'], 'CÁLCULOS': COR['mudo'], 'SIM_DÍVIDAS': COR['mudo'],
              'BACKUP': COR['mudo']}


def fnt(size=11, bold=False, italic=False, color=None, underline=None):
    return Font(name='Calibri', size=size, bold=bold, italic=italic, color=color or COR['texto'],
                underline=underline)


def fill(hexc):
    return PatternFill('solid', start_color=hexc, end_color=hexc)


thin = Side(style='thin', color=COR['borda'])
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def absref(ref):
    """B5:C9 -> $B$5:$C$9"""
    return re.sub(r'([A-Z]+)(\d+)', r'$\1$\2', ref)


class Ctx:
    def __init__(self, wb, sample=True):
        self.wb = wb
        self.sample = sample
        self.ws = {}
        self.clear_ranges = []   # (aba, intervalo, descrição) para a instrução de limpeza
        self.refs = {}

    def name(self, nome, sheet, ref):
        self.refs[nome] = (sheet, ref.replace('$', ''))
        self.wb.defined_names[nome] = DefinedName(nome, attr_text=f"'{sheet}'!{absref(ref)}")

    def name_formula(self, nome, formula):
        self.wb.defined_names[nome] = DefinedName(nome, attr_text=formula)


KINDS = {
    'in':   dict(fill=COR['azulclaro'], border=True, locked=False),
    'calc': dict(fill=COR['cinza'], border=True, locked=True),
    'hdr':  dict(fill=COR['header'], border=True, locked=True, color='FFFFFF', bold=True, h='center', wrap=True),
    'lbl':  dict(locked=True),
    'lblb': dict(locked=True, bold=True),
    'note': dict(locked=True, italic=True, color=COR['mudo']),
    'sec':  dict(locked=True, bold=True, size=13, color=COR['header']),
    'plain': dict(locked=True),
}


class Sh:
    def __init__(self, ctx, title):
        self.ctx = ctx
        self.title = title
        self.ws = ctx.ws[title]

    # ---- escrita -------------------------------------------------
    def put(self, ref, v=None, kind='plain', fmt=None, bold=None, color=None, fillc=None, h=None,
            wrap=None, size=None, italic=None, border=None, v_align='center', unlocked=None):
        c = self.ws[ref]
        k = KINDS[kind]
        if v is not None:
            c.value = v
        c.font = fnt(size or k.get('size', 11), bold if bold is not None else k.get('bold', False),
                     italic if italic is not None else k.get('italic', False), color or k.get('color'))
        f = fillc or k.get('fill')
        if f:
            c.fill = fill(f)
        if border if border is not None else k.get('border'):
            c.border = BORDER
        c.alignment = Alignment(horizontal=h or k.get('h'), vertical=v_align,
                                wrap_text=wrap if wrap is not None else k.get('wrap', False))
        lock = k.get('locked', True) if unlocked is None else (not unlocked)
        c.protection = Protection(locked=lock)
        if fmt:
            c.number_format = fmt
        return c

    def block(self, ref, kind, fmt=None, **kw):
        """Aplica estilo a um intervalo (sem valores)."""
        for row in self.ws[ref]:
            for c in (row if isinstance(row, tuple) else (row,)):
                self.put(c.coordinate, None, kind, fmt, **kw)

    def widths(self, d):
        for k, v in d.items():
            self.ws.column_dimensions[k].width = v

    def merge(self, ref):
        self.ws.merge_cells(ref)

    def dv(self, ref, formula, allow_blank=True, kind='list', msg=None, **kw):
        d = DataValidation(type=kind, formula1=formula, allow_blank=allow_blank, showErrorMessage=True,
                           errorTitle='Valor inválido', error=msg or 'Escolha um valor da lista.', **kw)
        self.ws.add_data_validation(d)
        d.add(ref)
        return d

    def cf(self, ref, rule):
        self.ws.conditional_formatting.add(ref, rule)

    def cf_formula(self, ref, formula, fillc=None, color=None, bold=None):
        kw = {}
        if fillc:
            kw['fill'] = PatternFill(start_color=fillc, end_color=fillc, fill_type='solid')
        if color or bold:
            kw['font'] = Font(color=color, bold=bold)
        self.ws.conditional_formatting.add(ref, FormulaRule(formula=[formula], stopIfTrue=False, **kw))

    # ---- cabeçalho padrão ---------------------------------------
    def header(self, titulo, desc, tab=None, ncols=14):
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 2
        ws.sheet_properties.tabColor = tab or TAB_COLORS.get(self.title, COR['azul'])
        self.put('B1', titulo, 'plain', size=20, bold=True, color=COR['header'])
        ws.row_dimensions[1].height = 30
        self.put('B2', desc, 'plain', italic=True, color=COR['mudo'], wrap=True, v_align='top')
        last = CL(max(ncols, 6) + 1)
        ws.merge_cells(f'B2:{last}2')
        ws.row_dimensions[2].height = 32
        # menu de navegação em 2 linhas
        items = [n for n in NAV if n[1] != self.title]
        per_row = 9
        col = 2
        for i, (lbl, tgt) in enumerate(items):
            r = 3 + i // per_row
            c = 2 + (i % per_row)
            ref = f'{CL(c)}{r}'
            cell = ws[ref]
            cell.value = lbl
            cell.hyperlink = Hyperlink(ref=ref, location=f"'{tgt}'!A1", display=lbl)
            cell.font = fnt(10, bold=(i == 0), color=(COR['verde'] if False else COR['azul']), underline='single')
            cell.alignment = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
            cell.fill = fill(COR['azul_l'] if i else COR['verde_l'])
            cell.border = BORDER
            cell.protection = Protection(locked=True)
        ws.row_dimensions[3].height = 18
        ws.row_dimensions[4].height = 18
        # largura mínima para o menu
        for c in range(2, 11):
            cd = ws.column_dimensions[CL(c)]
            if (cd.width or 0) < 12:
                cd.width = 12

    def finish(self, freeze=None, protect=True, landscape=True, allow_filter=False):
        ws = self.ws
        if freeze:
            ws.freeze_panes = freeze
        ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
        ws.page_setup.paperSize = 9
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.5
        if protect:
            ws.protection.sheet = True
            ws.protection.formatColumns = False
            ws.protection.formatRows = False
            if allow_filter:
                ws.protection.autoFilter = False
                ws.protection.sort = False


def mm_label(idx_expr):
    """Texto mm/aaaa a partir de um índice de mês (ano*12+mês), sem TEXT() (independe de idioma)."""
    return (f'RIGHT("0"&(MOD(({idx_expr})-1,12)+1),2)&"/"&INT((({idx_expr})-1)/12)')


def date_label(d):
    """dd/mm/aaaa sem TEXT()."""
    return f'RIGHT("0"&DAY({d}),2)&"/"&RIGHT("0"&MONTH({d}),2)&"/"&YEAR({d})'


def my_label(d):
    return f'RIGHT("0"&MONTH({d}),2)&"/"&YEAR({d})'


def SM(tipo, idx, extra=''):
    """Soma de LANÇAMENTOS (valor da parcela) de um tipo cujo período cobre o índice de mês."""
    return (f'SUMIFS(L_VParc,L_Tipo,"{tipo}"{extra},L_IIni,"<="&({idx}),L_IFim,">="&({idx}))')


def ml_label(idx_expr):
    """Rótulo curto de mês (Jan/26) a partir de um índice de mês — sem ambiguidade de data."""
    return (f'INDEX(Lista_MesesAb,MOD(({idx_expr})-1,12)+1)&"/"&RIGHT(INT((({idx_expr})-1)/12),2)')
