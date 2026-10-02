"""Gera a Planilha Financeira Avançada.  Uso: python3 build.py [--limpa]"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import Workbook
from common import *
import part_base, part_calc, part_budget, part_metas, part_cards_debts, part_invest, part_proj, part_dash, sample


def main(use_sample=True, out='../Planilha_Financeira_Avancada.xlsx', only=None):
    wb = Workbook()
    wb.remove(wb.active)
    ctx = Ctx(wb, use_sample)
    for n in SHEETS:
        ctx.ws[n] = wb.create_sheet(n)
    part_base.build_tabelas(ctx)
    part_base.build_config(ctx)
    part_base.build_lancamentos(ctx)
    part_base.build_recorrentes(ctx)
    part_calc.build_calc_core(ctx)
    part_metas.build_metas(ctx)
    part_metas.build_hist24(ctx)
    part_budget.build_orcamento(ctx)
    part_budget.build_fluxo(ctx)
    part_cards_debts.build_cartoes(ctx)
    part_cards_debts.build_sim_dividas(ctx)
    part_cards_debts.build_dividas(ctx)
    part_invest.build_invest(ctx)
    part_invest.build_patrimonio(ctx)
    part_invest.build_ir(ctx)
    part_proj.build_calc_extra(ctx)
    part_proj.build_proj(ctx)
    part_proj.build_calendario(ctx)
    part_proj.build_analises(ctx)
    part_dash.build_kpis(ctx)
    part_dash.build_dashboard(ctx)
    part_dash.build_inicio(ctx)
    part_dash.build_backup(ctx)
    if use_sample:
        sample.fill_sample(ctx)
    ctx.calc.sh.finish(protect=True)
    for n in ('TABELAS', 'CÁLCULOS', 'SIM_DÍVIDAS'):
        ctx.ws[n].sheet_state = 'hidden'
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    wb.active = wb.sheetnames.index('DASHBOARD')
    wb['DASHBOARD'].sheet_view.tabSelected = True
    wb.calculation.fullCalcOnLoad = True
    if only:
        for ws in wb.worksheets:
            if ws.title not in only:
                ws.sheet_state = 'hidden'
                from openpyxl.worksheet.print_settings import PrintArea
                ws._print_area = PrintArea()
        wb.active = wb.sheetnames.index(only[0])
        for ws in wb.worksheets:
            ws.sheet_view.tabSelected = (ws.title == only[0])
    wb.save(out)
    return ctx


if __name__ == '__main__':
    if '--render' in sys.argv:
        i = sys.argv.index('--render')
        main(True, '../.teste/render.xlsx', only=sys.argv[i + 1].split(','))
    elif '--limpa' in sys.argv:
        main(False, '../Planilha_Financeira_Avancada_LIMPA.xlsx')
    else:
        main(True)
