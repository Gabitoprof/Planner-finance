"""TABELAS (listas), CONFIGURAÇÕES, LANÇAMENTOS e RECORRENTES."""
from openpyxl.worksheet.table import Table, TableStyleInfo
from common import *

MESES = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho', 'Julho', 'Agosto', 'Setembro', 'Outubro',
         'Novembro', 'Dezembro']
MESES_AB = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
TIPOS_INV = ['Renda fixa', 'Tesouro Direto', 'Ações', 'FIIs', 'ETFs', 'Cripto', 'Previdência', 'Exterior', 'Outros']
CAT_META = ['Viagem', 'Carro', 'Casa', 'Reserva', 'Quitar dívida', 'Aposentadoria', 'Estudos', 'Outros']
META_COR = {'Viagem': COR['azul'], 'Carro': COR['amarelo'], 'Casa': COR['verde'], 'Reserva': COR['roxo'],
            'Quitar dívida': COR['vermelho'], 'Aposentadoria': '6C7A89', 'Estudos': '1ABC9C', 'Outros': 'A0A7B0'}

LISTAS = {
    'Meses': MESES, 'MesesAb': MESES_AB, 'Anos': list(range(2020, 2041)),
    'Tipos': ['Receita', 'Despesa', 'Transferência', 'Investimento'],
    'Status': ['Pago', 'Pendente', 'Agendado'],
    'Prioridade': ['Alta', 'Média', 'Baixa'],
    'CatMeta': CAT_META,
    'TipoInv': TIPOS_INV,
    'Liquidez': ['Diária', 'D+1', 'D+30', 'No vencimento', 'Baixa'],
    'TipoConta': ['Conta corrente', 'Poupança', 'Carteira', 'Conta investimento'],
    'TipoDivida': ['Empréstimo', 'Financiamento veículo', 'Financiamento imobiliário', 'Cartão/Fatura parcelada',
                   'Cheque especial', 'Outros'],
    'StatusDivida': ['Em dia', 'Atrasada', 'Negociando', 'Quitada'],
    'Sistema': ['Price', 'SAC', 'Outro'],
    'SimNao': ['Sim', 'Não'],
    'Freq': ['Mensal', 'Anual'],
    'TipoRec': ['Receita', 'Despesa', 'Investimento'],
    'Natureza': ['Fixa', 'Variável'],
    'Essencial': ['Essencial', 'Não essencial'],
    'Regra': ['Necessidades', 'Desejos', 'Poupança'],
    'DiaSem': ['Domingo', 'Segunda', 'Terça', 'Quarta', 'Quinta', 'Sexta', 'Sábado'],
    'StatusIR': ['Feito', 'Pendente', 'Não se aplica'],
}

# capacidades dos cadastros
N_CONTAS, N_CART, N_CATD, N_CATR, N_SUB, N_FORMA, N_TAG = 10, 8, 18, 9, 40, 10, 15
N_METAS = 15
N_REC = 40


def dyn(sheet, first_col, r0, r1):
    s = f"'{sheet}'"
    c = first_col
    return (f"{s}!${c}${r0}:INDEX({s}!${c}${r0}:${c}${r1},MAX(1,COUNTA({s}!${c}${r0}:${c}${r1})))")


# =====================================================================
def build_tabelas(ctx):
    sh = Sh(ctx, 'TABELAS')
    ws = sh.ws
    sh.header('TABELAS (apoio)', 'Listas auxiliares usadas nas listas suspensas. Não precisa mexer aqui. '
              'Esta aba pode ficar oculta (botão direito na aba > Reexibir).', ncols=12)
    col = 2
    ctx.tab_cols = {}
    for nome, vals in LISTAS.items():
        L = CL(col)
        sh.put(f'{L}6', nome, 'hdr')
        for i, v in enumerate(vals):
            sh.put(f'{L}{7 + i}', v, 'calc')
        ctx.name(f'Lista_{nome}', 'TABELAS', f'{L}7:{L}{6 + len(vals)}')
        ws.column_dimensions[L].width = 16
        ctx.tab_cols[nome] = L
        col += 1
    # listas combinadas (compactadas por fórmula)
    LC = CL(col + 1)
    sh.put(f'{LC}6', 'Conta ou Cartão', 'hdr')
    for k in range(1, N_CONTAS + N_CART + 1):
        r = 6 + k
        sh.put(f'{LC}{r}',
               f'=IF({k}<=COUNTA(Contas_Nome),INDEX(Contas_Nome,{k}),IF({k}<=COUNTA(Contas_Nome)+COUNTA(Cartoes_Nome),'
               f'INDEX(Cartoes_Nome,{k}-COUNTA(Contas_Nome)),""))', 'calc')
    r1 = 6 + N_CONTAS + N_CART
    ctx.name_formula('Lista_ContaCartao',
                     f"'TABELAS'!${LC}$7:INDEX('TABELAS'!${LC}$7:${LC}${r1},"
                     f"MAX(1,COUNTA(Contas_Nome)+COUNTA(Cartoes_Nome)))")
    ws.column_dimensions[LC].width = 28
    LK = CL(col + 2)
    sh.put(f'{LK}6', 'Categorias (todas)', 'hdr')
    ntot = N_CATD + N_CATR + 2
    for k in range(1, ntot + 1):
        r = 6 + k
        sh.put(f'{LK}{r}',
               f'=IF({k}<=COUNTA(CatD_Nome),INDEX(CatD_Nome,{k}),IF({k}<=COUNTA(CatD_Nome)+COUNTA(CatR_Nome),'
               f'INDEX(CatR_Nome,{k}-COUNTA(CatD_Nome)),IF({k}=COUNTA(CatD_Nome)+COUNTA(CatR_Nome)+1,"Investimento",'
               f'IF({k}=COUNTA(CatD_Nome)+COUNTA(CatR_Nome)+2,"Transferência",""))))', 'calc')
    ctx.name_formula('Lista_Categorias',
                     f"'TABELAS'!${LK}$7:INDEX('TABELAS'!${LK}$7:${LK}${6 + ntot},"
                     f"MAX(1,COUNTA(CatD_Nome)+COUNTA(CatR_Nome)+2))")
    ws.column_dimensions[LK].width = 28
    # cores das categorias de metas (usadas nas formatações condicionais)
    LM = CL(col + 4)
    sh.put(f'{LM}6', 'Cor da meta', 'hdr')
    sh.put(f'{CL(col + 5)}6', 'Hex', 'hdr')
    for i, c in enumerate(CAT_META):
        sh.put(f'{LM}{7 + i}', c, 'calc', fillc=META_COR[c], color='FFFFFF', bold=True)
        sh.put(f'{CL(col + 5)}{7 + i}', META_COR[c], 'calc')
    ws.column_dimensions[LM].width = 18
    sh.finish(freeze='B7')
    return sh


# =====================================================================
def build_config(ctx):
    sh = Sh(ctx, 'CONFIGURAÇÕES')
    ws = sh.ws
    sh.header('CONFIGURAÇÕES', 'Preencha uma vez: dados gerais, metas e os cadastros (contas, cartões, categorias...). '
              'Células azuis = você digita. Preencha as listas de cima para baixo, sem deixar linhas em branco no meio.',
              ncols=8)
    sh.widths({'B': 40, 'C': 22, 'D': 20, 'E': 20, 'F': 24, 'G': 40})
    sh.put('B5', '⬇ digite nas células azuis', 'note', color=COR['azul'], bold=True)

    def row_param(r, label, nm, fmt=None, note='', dv=None):
        sh.put(f'B{r}', label, 'lbl')
        sh.put(f'C{r}', None, 'in', fmt, h='right')
        sh.put(f'D{r}', note, 'note')
        ctx.name(nm, 'CONFIGURAÇÕES', f'C{r}')
        if dv:
            sh.dv(f'C{r}', dv)

    sh.put('B6', '1. Dados gerais', 'sec')
    row_param(7, 'Nome', 'Cfg_Nome', note='Seu nome (aparece no Dashboard).')
    row_param(8, 'Ano de referência', 'Cfg_AnoIn', INT, 'Deixe vazio para usar o ano de hoje.')
    row_param(9, 'Mês atual', 'Cfg_MesIn', None, 'Deixe vazio para usar o mês de hoje.', '=Lista_Meses')
    row_param(10, 'Renda mensal líquida média (R$)', 'Cfg_Renda', BRLZ, 'Quanto cai na conta por mês, em média.')
    row_param(11, 'Meses da reserva de emergência', 'Cfg_MesesReserva', INT, 'Padrão: 6 meses de despesas.')
    row_param(12, '% da meta de poupança (sobre a renda)', 'Cfg_MetaPoup', PCT, 'Padrão: 20%.')
    row_param(13, 'Inflação anual estimada', 'Cfg_Infl', '0.00%', 'Usada nas projeções e na comparação com investimentos.')
    row_param(14, 'Taxa Selic/CDI anual de referência', 'Cfg_CDI', '0.00%', 'Atualize quando mudar (apenas referência).')
    row_param(15, 'Data de hoje', 'Cfg_Hoje', DATA,
              'Usa a data do computador (=HOJE()). Pode digitar uma data fixa para testar.')
    ws['C15'] = '=TODAY()'
    row_param(16, 'Orçamento base zero?', 'Cfg_BaseZero', None,
              'Sim = toda a receita planejada deve ser alocada (veja ORÇAMENTO).', '=Lista_SimNao')
    row_param(17, 'Alerta de endividamento acima de', 'Cfg_LimDiv', PCT, 'Parcelas de dívidas / renda.')
    row_param(18, 'Alerta de uso do limite do cartão acima de', 'Cfg_LimCartao', PCT, 'Padrão: 70%.')
    sh.put('B20', '2. Metas mensais de comportamento', 'sec')
    row_param(21, 'Meta de gasto máximo total (R$/mês)', 'Cfg_MaxGasto', BRLZ,
              'Metas por categoria: coluna "Gasto máximo" na seção 5 abaixo.')
    row_param(22, 'Meta de aporte mensal (R$)', 'Cfg_MetaAporte', BRLZ, 'Quanto pretende investir/guardar por mês.')
    row_param(23, 'Meta de patrimônio líquido ao fim do ano (R$)', 'Cfg_MetaPatr', BRLZ, '')
    ctx.cfg_params = {}
    for ref, v in (('C11', 6), ('C12', 0.2), ('C17', 0.3), ('C18', 0.7)):
        ws[ref] = v

    # --- blocos de cadastro --------------------------------------
    def cadastro(r, titulo, heads, n, names, fmts=None, dvs=None, extra_calc=None, note=None):
        sh.put(f'B{r}', titulo, 'sec')
        if note:
            sh.put(f'D{r}', note, 'note')
        for i, h in enumerate(heads):
            sh.put(f'{CL(2 + i)}{r + 1}', h, 'hdr')
        ws.row_dimensions[r + 1].height = 32
        r0, r1 = r + 2, r + 1 + n
        for i in range(len(heads)):
            L = CL(2 + i)
            kind = 'calc' if (extra_calc and i == extra_calc[0]) else 'in'
            sh.block(f'{L}{r0}:{L}{r1}', kind, (fmts or {}).get(i))
            if names.get(i):
                ctx.name(names[i], 'CONFIGURAÇÕES', f'{L}{r0}:{L}{r1}')
            if dvs and i in dvs:
                sh.dv(f'{L}{r0}:{L}{r1}', dvs[i])
        return r0, r1

    r0, r1 = cadastro(25, '3. Contas bancárias', ['Conta', 'Tipo', 'Saldo inicial (R$)', 'Saldo atual (calculado)'],
                      N_CONTAS, {0: 'Contas_Nome', 1: 'Contas_Tipo', 2: 'Contas_Saldo0', 3: 'Contas_SaldoAtual'},
                      {2: BRLZ, 3: BRLZ}, {1: '=Lista_TipoConta'}, extra_calc=(3, None),
                      note='Saldo atual = saldo inicial + lançamentos PAGOS da conta.')
    for r in range(r0, r1 + 1):
        ws[f'E{r}'] = (f'=IF(B{r}="","",D{r}+SUMIFS(L_VParc,L_Conta,B{r},L_Tipo,"Receita",L_Status,"Pago")'
                       f'-SUMIFS(L_VParc,L_Conta,B{r},L_Tipo,"Despesa",L_Status,"Pago")'
                       f'-SUMIFS(L_VParc,L_Conta,B{r},L_Tipo,"Investimento",L_Status,"Pago")'
                       f'+SUMIFS(L_VParc,L_Conta,B{r},L_Tipo,"Transferência",L_Status,"Pago"))')
    ctx.sample_ranges = {'contas': (r0, r1)}
    ctx.dvname = {}

    r0c, r1c = cadastro(r1 + 2, '4. Cartões de crédito', ['Cartão', 'Limite (R$)', 'Dia de fechamento', 'Dia de vencimento'],
                        N_CART, {0: 'Cartoes_Nome', 1: 'Cartoes_Limite', 2: 'Cartoes_Fech', 3: 'Cartoes_Venc'},
                        {1: BRLZ, 2: INT, 3: INT},
                        note='Compras a partir do dia de fechamento entram na fatura seguinte.')
    sh.dv(f'D{r0c}:E{r1c}', '1', kind='whole', operator='greaterThanOrEqual', msg='Informe um dia de 1 a 28.')
    ws.data_validations.dataValidation[-1].formula2 = '28'
    ws.data_validations.dataValidation[-1].operator = 'between'
    ctx.sample_ranges['cartoes'] = (r0c, r1c)

    r0d, r1d = cadastro(r1c + 2, '5. Categorias de despesa',
                        ['Categoria', 'Fixa / Variável', 'Essencial / Não essencial', 'Regra 50/30/20',
                         'Gasto máximo mensal (R$)'],
                        N_CATD, {0: 'CatD_Nome', 1: 'CatD_Nat', 2: 'CatD_Ess', 3: 'CatD_Regra', 4: 'CatD_Max'},
                        {4: BRLZ}, {1: '=Lista_Natureza', 2: '=Lista_Essencial', 3: '=Lista_Regra'},
                        note='Para criar uma categoria nova: escreva na 1ª linha vazia.')
    ctx.sample_ranges['catd'] = (r0d, r1d)
    r0r, r1r = cadastro(r1d + 2, '6. Categorias de receita', ['Categoria'], N_CATR, {0: 'CatR_Nome'})
    ctx.sample_ranges['catr'] = (r0r, r1r)
    r0s, r1s = cadastro(r1r + 2, '7. Subcategorias', ['Subcategoria', 'Categoria (pai)'], N_SUB,
                        {0: 'Sub_Nome', 1: 'Sub_Cat'}, None, {1: '=Lista_Categorias'})
    ctx.sample_ranges['sub'] = (r0s, r1s)
    r0f, r1f = cadastro(r1s + 2, '8. Formas de pagamento', ['Forma de pagamento'], N_FORMA, {0: 'Forma_Nome'})
    ctx.sample_ranges['formas'] = (r0f, r1f)
    r0t, r1t = cadastro(r1f + 2, '9. Tags (ex.: #viagem, #presente)', ['Tag'], N_TAG, {0: 'Tag_Nome'})
    ctx.sample_ranges['tags'] = (r0t, r1t)

    # listas dinâmicas (sem linhas vazias)
    S = 'CONFIGURAÇÕES'
    ctx.name_formula('Lista_Contas', dyn(S, 'B', *ctx.sample_ranges['contas']))
    ctx.name_formula('Lista_Cartoes', dyn(S, 'B', *ctx.sample_ranges['cartoes']))
    ctx.name_formula('Lista_CatD', dyn(S, 'B', *ctx.sample_ranges['catd']))
    ctx.name_formula('Lista_CatR', dyn(S, 'B', *ctx.sample_ranges['catr']))
    ctx.name_formula('Lista_Subcats', dyn(S, 'B', *ctx.sample_ranges['sub']))
    ctx.name_formula('Lista_Formas', dyn(S, 'B', *ctx.sample_ranges['formas']))
    ctx.name_formula('Lista_Tags', dyn(S, 'B', *ctx.sample_ranges['tags']))

    # preenchimento padrão (cadastros base que valem também para a versão limpa)
    cat_d = [('Moradia', 'Fixa', 'Essencial', 'Necessidades'), ('Alimentação', 'Variável', 'Essencial', 'Necessidades'),
             ('Transporte', 'Variável', 'Essencial', 'Necessidades'), ('Saúde', 'Variável', 'Essencial', 'Necessidades'),
             ('Educação', 'Fixa', 'Essencial', 'Necessidades'), ('Lazer', 'Variável', 'Não essencial', 'Desejos'),
             ('Assinaturas', 'Fixa', 'Não essencial', 'Desejos'), ('Vestuário', 'Variável', 'Não essencial', 'Desejos'),
             ('Dívidas e Financiamentos', 'Fixa', 'Essencial', 'Necessidades'),
             ('Impostos e Taxas', 'Fixa', 'Essencial', 'Necessidades'),
             ('Família/Pets', 'Variável', 'Essencial', 'Necessidades'), ('Doações', 'Variável', 'Não essencial', 'Desejos'),
             ('Outros', 'Variável', 'Não essencial', 'Desejos')]
    for i, (n, a, b, c) in enumerate(cat_d):
        r = r0d + i
        ws[f'B{r}'], ws[f'C{r}'], ws[f'D{r}'], ws[f'E{r}'] = n, a, b, c
    for i, n in enumerate(['Salário', '13º/Férias/PLR', 'Freelance/Extra', 'Aluguel recebido', 'Rendimentos',
                           'Reembolsos', 'Outros']):
        ws[f'B{r0r + i}'] = n
    for i, n in enumerate(['Pix', 'Débito', 'Crédito', 'Dinheiro', 'Boleto', 'Débito automático', 'Transferência']):
        ws[f'B{r0f + i}'] = n
    for i, n in enumerate(['#viagem', '#presente', '#reforma', '#natal', '#trabalho']):
        ws[f'B{r0t + i}'] = n
    sh.finish(freeze='A6')
    return sh


# =====================================================================
LCOLS = [('Data', 'A', 12, DATA), ('Descrição', 'B', 32, None), ('Tipo', 'C', 14, None), ('Categoria', 'D', 22, None),
         ('Subcategoria', 'E', 18, None), ('Conta ou Cartão', 'F', 24, None), ('Forma de pagamento', 'G', 18, None),
         ('Parcela atual', 'H', 9, INT), ('Total de parcelas', 'I', 9, INT), ('Status', 'J', 11, None),
         ('Tag', 'K', 12, None), ('Meta vinculada', 'L', 22, None), ('Observação', 'M', 22, None),
         ('Valor', 'N', 14, BRLZ),
         ('Mês', 'O', 6, INT), ('Ano', 'P', 7, INT), ('Competência', 'Q', 11, MESANO), ('Valor da parcela', 'R', 14, BRLZ),
         ('Fatura de referência', 'S', 11, MESANO), ('Alerta', 'T', 30, None),
         ('Idx início', 'U', 9, INT), ('Idx fim', 'V', 9, INT), ('Fatura início', 'W', 9, INT), ('Fatura fim', 'X', 9, INT),
         ('Parcelas a vencer (cartão)', 'Y', 12, BRLZ), ('Dia', 'Z', 5, INT), ('Dia da semana', 'AA', 7, INT),
         ('Gasto no mês sel.', 'AB', 12, BRLZ), ('Chave calendário', 'AC', 10, INT), ('Ordem calendário', 'AD', 8, INT)]
LNAMES = {'A': 'L_Data', 'B': 'L_Desc', 'C': 'L_Tipo', 'D': 'L_Cat', 'E': 'L_Sub', 'F': 'L_Conta', 'G': 'L_Forma',
          'H': 'L_ParcAt', 'I': 'L_ParcTot', 'J': 'L_Status', 'K': 'L_Tag', 'L': 'L_Meta', 'M': 'L_Obs', 'N': 'L_Valor',
          'O': 'L_Mes', 'P': 'L_Ano', 'Q': 'L_Comp', 'R': 'L_VParc', 'S': 'L_Fat', 'T': 'L_Alerta', 'U': 'L_IIni',
          'V': 'L_IFim', 'W': 'L_FIni', 'X': 'L_FFim', 'Y': 'L_Aberto', 'Z': 'L_Dia', 'AA': 'L_Sem', 'AB': 'L_DespSel',
          'AC': 'L_CalChave', 'AD': 'L_CalOrdem'}


def lanc_formula(col, r):
    f = {
        'O': f'=IF($A{r}="","",MONTH($A{r}))',
        'P': f'=IF($A{r}="","",YEAR($A{r}))',
        'Q': f'=IF($A{r}="","",DATE(P{r},O{r},1))',
        'R': f'=IF(OR($A{r}="",$N{r}=""),"",IF(AND($C{r}<>"Transferência",N($I{r})>1),$N{r}/$I{r},$N{r}))',
        'S': (f'=IF(OR($A{r}="",$F{r}="",$C{r}<>"Despesa"),"",IF(COUNTIF(Cartoes_Nome,$F{r})=0,"",'
              f'DATE(P{r},O{r}+IF(DAY($A{r})>=INDEX(Cartoes_Fech,MATCH($F{r},Cartoes_Nome,0)),1,0),1)))'),
        'T': (f'=IF($A{r}="","",IF(COUNTIFS(L_Data,$A{r},L_Desc,$B{r},L_Valor,$N{r},L_Tipo,$C{r})>1,'
              f'"⚠ Possível duplicidade",IF(OR(AND($C{r}="Despesa",COUNTIF(CatD_Nome,$D{r})=0),'
              f'AND($C{r}="Receita",COUNTIF(CatR_Nome,$D{r})=0)),"⚠ Categoria não combina com o tipo",'
              f'IF(OR($N{r}="",$C{r}=""),"⚠ Falta valor ou tipo",""))))'),
        'U': f'=IF($A{r}="",0,P{r}*12+O{r})',
        'V': f'=IF($A{r}="",0,U{r}+MAX(0,N($I{r})-MAX(N($H{r}),1)))',
        'W': f'=IF(S{r}="",0,YEAR(S{r})*12+MONTH(S{r}))',
        'X': f'=IF(W{r}=0,0,W{r}+V{r}-U{r})',
        'Y': (f'=IF(W{r}=0,0,N(R{r})*MAX(0,X{r}-MAX(W{r},YEAR(Cfg_Hoje)*12+MONTH(Cfg_Hoje)+'
              f'IF(DAY(Cfg_Hoje)>=INDEX(Cartoes_Fech,MATCH($F{r},Cartoes_Nome,0)),1,0))+1))'),
        'Z': f'=IF($A{r}="",0,DAY($A{r}))',
        'AA': f'=IF($A{r}="",0,WEEKDAY($A{r}))',
        'AB': f'=IF(AND($C{r}="Despesa",U{r}>0,U{r}<=Sel_Idx,V{r}>=Sel_Idx),N(R{r})+ROW()/1000000000,0)',
        'AC': f'=IF(AND($C{r}="Despesa",U{r}>0,U{r}<=Sel_Idx,V{r}>=Sel_Idx),Z{r}*10000+ROW(),0)',
        'AD': f'=IF(AC{r}=0,0,COUNTIFS(L_CalChave,">0",L_CalChave,"<"&AC{r})+1)',
    }
    return f[col]


def build_lancamentos(ctx):
    sh = Sh(ctx, 'LANÇAMENTOS')
    ws = sh.ws
    sh.header('LANÇAMENTOS', 'A ÚNICA aba de digitação diária: uma linha por receita, despesa, transferência ou aporte. '
              'Compra parcelada = UMA linha (valor TOTAL, parcela atual e total de parcelas). '
              'Pagamento de fatura do cartão: lance como Transferência (valor negativo) na conta.', ncols=14)
    sh.put('B5', '⬇ digite nas colunas azuis (A a N)', 'note', color=COR['azul'], bold=True)
    sh.put('O5', 'calculado automaticamente ⬇', 'note')
    hr = LR0 - 1
    for head, L, w, fm in LCOLS:
        sh.put(f'{L}{hr}', head, 'hdr')
        ws.column_dimensions[L].width = w
    ws.row_dimensions[hr].height = 34
    # estilos das colunas
    for head, L, w, fm in LCOLS:
        kind = 'in' if CI(L) <= 14 else 'calc'
        for r in range(LR0, LR1 + 1):
            c = ws[f'{L}{r}']
            if kind == 'calc':
                c.value = lanc_formula(L, r)
            sh.put(f'{L}{r}', None, kind, fm, h=('left' if L in ('B', 'M', 'T') else None))
        ctx.name(LNAMES[L], 'LANÇAMENTOS', f'{L}{LR0}:{L}{LR1}')
    for L in ['U', 'V', 'W', 'X', 'Y', 'Z', 'AA', 'AB', 'AC', 'AD']:
        ws.column_dimensions[L].hidden = True
    rng = lambda L: f'{L}{LR0}:{L}{LR1}'
    sh.dv(rng('A'), '36526', kind='date', operator='greaterThanOrEqual', msg='Digite uma data válida (dd/mm/aaaa).')
    sh.dv(rng('C'), '=Lista_Tipos')
    sh.dv(rng('D'), '=Lista_Categorias')
    sh.dv(rng('E'), '=Lista_Subcats')
    sh.dv(rng('F'), '=Lista_ContaCartao')
    sh.dv(rng('G'), '=Lista_Formas')
    sh.dv(rng('J'), '=Lista_Status')
    sh.dv(rng('K'), '=Lista_Tags')
    sh.dv(rng('L'), '=Lista_MetasNomes')
    sh.dv(rng('H'), '1', kind='whole', operator='greaterThanOrEqual', msg='Informe um número inteiro (1, 2, 3...).')
    sh.dv(rng('I'), '1', kind='whole', operator='greaterThanOrEqual', msg='Informe um número inteiro (1, 2, 3...).')
    # formatação condicional
    sh.cf_formula(rng('J'), f'$J{LR0}="Pago"', COR['verde_l'], COR['verde'], True)
    sh.cf_formula(rng('J'), f'$J{LR0}="Pendente"', COR['amarelo_l'], '9A5B00', True)
    sh.cf_formula(rng('J'), f'$J{LR0}="Agendado"', COR['azul_l'], COR['azul'], True)
    sh.cf_formula(rng('C'), f'$C{LR0}="Receita"', None, COR['verde'], True)
    sh.cf_formula(rng('C'), f'$C{LR0}="Despesa"', None, COR['vermelho'], True)
    sh.cf_formula(rng('C'), f'$C{LR0}="Investimento"', None, COR['azul'], True)
    sh.cf_formula(rng('T'), f'LEFT($T{LR0},1)="⚠"', COR['vermelho_l'], COR['vermelho'], True)
    sh.cf_formula(f'N{LR0}:N{LR1}', f'$C{LR0}="Receita"', None, COR['verde'])
    sh.cf_formula(f'N{LR0}:N{LR1}', f'$C{LR0}="Despesa"', None, COR['vermelho'])
    tab = Table(displayName='tblLancamentos', ref=f'A{hr}:AD{LR1}')
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False, showColumnStripes=False)
    ws.add_table(tab)
    ws.auto_filter.ref = None
    ws.print_title_rows = f'{hr}:{hr}'
    ws.print_area = f'A1:T{LR1}'
    sh.finish(freeze=f'C{LR0}', allow_filter=True)
    ctx.clear_ranges.append(('LANÇAMENTOS', f'A{LR0}:N{LR1}', 'Lançamentos de exemplo'))
    ctx.name('Exemplo_Lancamentos', 'LANÇAMENTOS', f'A{LR0}:N{LR1}')
    return sh


# =====================================================================
RC = [('Descrição', 'B', 30), ('Tipo', 'C', 14), ('Categoria', 'D', 22), ('Conta ou Cartão', 'E', 24),
      ('Valor (R$)', 'F', 14), ('Frequência', 'G', 12), ('Dia do mês', 'H', 9), ('Início', 'I', 12),
      ('Fim (opcional)', 'J', 12), ('Ativo?', 'K', 9),
      ('Idx início', 'L', 9), ('Idx fim', 'M', 9), ('Mês de ref.', 'N', 8), ('Valor anual', 'O', 14),
      ('É assinatura?', 'P', 9), ('Ordem assinatura', 'Q', 9)]
RNAMES = {'B': 'Rec_Desc', 'C': 'Rec_Tipo', 'D': 'Rec_Cat', 'E': 'Rec_Conta', 'F': 'Rec_Valor', 'G': 'Rec_Freq',
          'H': 'Rec_Dia', 'I': 'Rec_Ini', 'J': 'Rec_Fim', 'K': 'Rec_Ativo', 'L': 'Rec_IIni', 'M': 'Rec_IFim',
          'N': 'Rec_MesRef', 'O': 'Rec_Anual', 'P': 'Rec_Assin', 'Q': 'Rec_AssinOrd'}
RR0 = 7
RR1 = RR0 + N_REC - 1


def build_recorrentes(ctx):
    sh = Sh(ctx, 'RECORRENTES')
    ws = sh.ws
    sh.header('RECORRENTES', 'Cadastre salário, aluguel, assinaturas e outros itens que se repetem. Eles entram '
              'automaticamente nas PROJEÇÕES dos meses futuros (não precisa lançar mês a mês no futuro). '
              'Não cadastre o mesmo item aqui e em LANÇAMENTOS futuros.', ncols=11)
    sh.put('B5', '⬇ digite nas colunas azuis (B a K)', 'note', color=COR['azul'], bold=True)
    for head, L, w in RC:
        sh.put(f'{L}6', head, 'hdr')
        ws.column_dimensions[L].width = w
    ws.row_dimensions[6].height = 30
    for head, L, w in RC:
        kind = 'in' if CI(L) <= 11 else 'calc'
        fm = {'F': BRLZ, 'I': DATA, 'J': DATA, 'O': BRL, 'H': INT}.get(L)
        for r in range(RR0, RR1 + 1):
            if kind == 'calc':
                ws[f'{L}{r}'] = {
                    'L': f'=IF(OR($B{r}="",NOT(ISNUMBER($I{r}))),0,YEAR($I{r})*12+MONTH($I{r}))',
                    'M': f'=IF(L{r}=0,0,IF(ISNUMBER($J{r}),YEAR($J{r})*12+MONTH($J{r}),999999))',
                    'N': f'=IF(L{r}=0,0,MONTH($I{r}))',
                    'O': f'=IF(L{r}=0,0,IF($G{r}="Anual",N($F{r}),N($F{r})*12))',
                    'P': f'=IF(AND($C{r}="Despesa",$D{r}="Assinaturas",$K{r}="Sim",L{r}>0),1,0)',
                    'Q': f'=IF(P{r}=1,SUM(P${RR0}:P{r}),0)',
                }[L]
            sh.put(f'{L}{r}', None, kind, fm)
        ctx.name(RNAMES[L], 'RECORRENTES', f'{L}{RR0}:{L}{RR1}')
    sh.dv(f'C{RR0}:C{RR1}', '=Lista_TipoRec')
    sh.dv(f'D{RR0}:D{RR1}', '=Lista_Categorias')
    sh.dv(f'E{RR0}:E{RR1}', '=Lista_ContaCartao')
    sh.dv(f'G{RR0}:G{RR1}', '=Lista_Freq')
    sh.dv(f'K{RR0}:K{RR1}', '=Lista_SimNao')
    sh.dv(f'H{RR0}:H{RR1}', '1', kind='whole', operator='between', msg='Dia de 1 a 31.')
    ws.data_validations.dataValidation[-1].formula2 = '31'
    for L in 'LMNPQ':
        ws.column_dimensions[L].hidden = True
    sh.cf_formula(f'K{RR0}:K{RR1}', f'$K{RR0}="Não"', COR['cinza'], COR['mudo'])
    tab = Table(displayName='tblRecorrentes', ref=f'B6:Q{RR1}')
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    ws.add_table(tab)
    ws.print_title_rows = '6:6'
    ws.print_area = f'A1:O{RR1}'
    sh.finish(freeze=f'C{RR0}', allow_filter=True)
    ctx.name('Exemplo_Recorrentes', 'RECORRENTES', f'B{RR0}:K{RR1}')
    ctx.clear_ranges.append(('RECORRENTES', f'B{RR0}:K{RR1}', 'Recorrentes de exemplo'))
    return sh


def REC(tipo, idx, extra=''):
    """Soma dos recorrentes ativos de um tipo no mês de índice idx (mensais + anuais no mês de início)."""
    base = f'Rec_Tipo,"{tipo}",Rec_Ativo,"Sim"{extra},Rec_IIni,"<="&({idx}),Rec_IFim,">="&({idx})'
    return (f'(SUMIFS(Rec_Valor,{base},Rec_Freq,"<>Anual")+'
            f'SUMIFS(Rec_Valor,{base},Rec_Freq,"Anual",Rec_MesRef,MOD(({idx})-1,12)+1))')
