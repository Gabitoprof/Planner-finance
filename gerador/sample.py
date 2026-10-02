"""Dados de EXEMPLO fictícios (todos marcados com 'EXEMPLO')."""
import datetime as dt
import random
from openpyxl.utils.cell import range_boundaries
from common import *
from part_base import LR0

D = dt.date
HOJE = D(2026, 10, 1)
MONTHS = [(2025, 9), (2025, 10), (2025, 11), (2025, 12)] + [(2026, k) for k in range(1, 11)]


def put_range(ctx, nome, rows):
    """Escreve uma matriz (lista de listas) a partir do canto superior esquerdo do nome definido."""
    sheet, ref = ctx.refs[nome]
    c1, r1, c2, r2 = range_boundaries(ref)
    ws = ctx.wb[sheet]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is not None:
                ws.cell(row=r1 + i, column=c1 + j, value=v)


def fill_config(ctx):
    ws = ctx.wb['CONFIGURAÇÕES']
    vals = {'Cfg_Nome': 'Fulano de Tal (EXEMPLO)', 'Cfg_AnoIn': 2026, 'Cfg_MesIn': 'Outubro', 'Cfg_Renda': 12800,
            'Cfg_Infl': 0.045, 'Cfg_CDI': 0.1375, 'Cfg_MaxGasto': 10500, 'Cfg_MetaAporte': 2100, 'Cfg_MetaPatr': 700000}
    for n, v in vals.items():
        sh, ref = ctx.refs[n]
        ctx.wb[sh][ref] = v
    put_range(ctx, 'Contas_Nome', [['Conta Corrente Banco Azul'], ['Poupança Banco Azul'], ['Carteira (dinheiro)']])
    put_range(ctx, 'Contas_Tipo', [['Conta corrente'], ['Poupança'], ['Carteira']])
    put_range(ctx, 'Contas_Saldo0', [[4200], [6000], [250]])
    put_range(ctx, 'Cartoes_Nome', [['Cartão Roxinho'], ['Cartão Azul Gold']])
    put_range(ctx, 'Cartoes_Limite', [[9000], [6000]])
    put_range(ctx, 'Cartoes_Fech', [[25], [10]])
    put_range(ctx, 'Cartoes_Venc', [[5], [17]])
    put_range(ctx, 'CatD_Max', [[3600], [1800], [900], [1100], [450], [700], [300], [400], [2200], [250], [450], [150], [300]])
    subs = [('Aluguel', 'Moradia'), ('Condomínio', 'Moradia'), ('Energia', 'Moradia'), ('Água', 'Moradia'), ('Internet', 'Moradia'),
            ('Mercado', 'Alimentação'), ('Restaurantes', 'Alimentação'), ('Delivery', 'Alimentação'),
            ('Combustível', 'Transporte'), ('App de transporte', 'Transporte'), ('Manutenção', 'Transporte'),
            ('Plano de saúde', 'Saúde'), ('Farmácia', 'Saúde'), ('Consultas', 'Saúde'), ('Cursos', 'Educação'),
            ('Cinema', 'Lazer'), ('Viagens', 'Lazer'), ('Bares', 'Lazer'), ('Streaming', 'Assinaturas'),
            ('Música', 'Assinaturas'), ('Nuvem', 'Assinaturas'), ('Academia', 'Assinaturas'), ('Roupas', 'Vestuário'),
            ('Financiamento', 'Dívidas e Financiamentos'), ('Empréstimo', 'Dívidas e Financiamentos'),
            ('IPTU', 'Impostos e Taxas'), ('IPVA', 'Impostos e Taxas'), ('Ração', 'Família/Pets'), ('Veterinário', 'Família/Pets'),
            ('ONG', 'Doações')]
    put_range(ctx, 'Sub_Nome', [[a, b] for a, b in subs])


# ------------------------------------------------------------------ lançamentos
def gerar_lancamentos():
    rnd = random.Random(2026)
    L = []     # (data, desc, tipo, cat, sub, conta, forma, parc_at, parc_tot, status, tag, meta, valor)
    CC, PO, CR, AZ = 'Conta Corrente Banco Azul', 'Poupança Banco Azul', 'Cartão Roxinho', 'Cartão Azul Gold'

    def st(d):
        return 'Pago' if d < HOJE else 'Agendado'

    def add(d, desc, tipo, cat, sub, conta, forma, valor, pa=None, pt=None, tag=None, meta=None, status=None):
        L.append([d, desc + ' (EXEMPLO)' if False else desc, tipo, cat, sub, conta, forma, pa, pt, status or st(d), tag, meta,
                  round(valor, 2)])

    for (y, m) in MONTHS:
        # receitas
        add(D(y, m, 5), 'Salário', 'Receita', 'Salário', None, CC, 'Transferência', 10800)
        if (y, m) == (2025, 12):
            add(D(y, m, 20), '13º salário', 'Receita', '13º/Férias/PLR', None, CC, 'Transferência', 10800)
        add(D(y, m, 10), 'Aluguel recebido - apto 12', 'Receita', 'Aluguel recebido', None, CC, 'Pix', 1900)
        if m in (2, 4, 6, 8):
            add(D(y, m, 18), 'Projeto freelance', 'Receita', 'Freelance/Extra', None, CC, 'Pix', rnd.choice([1200, 1800, 2400]))
        if (y, m) <= (2026, 9):
            add(D(y, m, 28), 'Rendimento da poupança', 'Receita', 'Rendimentos', None, PO, 'Transferência', round(rnd.uniform(60, 95), 2))
        if m == 5:
            add(D(y, m, 20), 'Reembolso plano de saúde', 'Receita', 'Reembolsos', None, CC, 'Pix', 340)
        # despesas fixas
        add(D(y, m, 10), 'Aluguel residencial', 'Despesa', 'Moradia', 'Aluguel', CC, 'Boleto', 2400)
        add(D(y, m, 10), 'Condomínio', 'Despesa', 'Moradia', 'Condomínio', CC, 'Boleto', 650)
        add(D(y, m, 15), 'Conta de energia', 'Despesa', 'Moradia', 'Energia', CC, 'Débito automático', round(rnd.uniform(140, 230), 2))
        add(D(y, m, 15), 'Conta de água', 'Despesa', 'Moradia', 'Água', CC, 'Débito automático', round(rnd.uniform(70, 110), 2))
        add(D(y, m, 20), 'Internet fibra', 'Despesa', 'Moradia', 'Internet', CC, 'Débito automático', 129.90)
        add(D(y, m, 5), 'Plano de saúde', 'Despesa', 'Saúde', 'Plano de saúde', CC, 'Débito automático', 780)
        add(D(y, m, 8), 'Curso de inglês', 'Despesa', 'Educação', 'Cursos', CC, 'Boleto', 390)
        add(D(y, m, 3), 'Academia', 'Despesa', 'Assinaturas', 'Academia', CC, 'Débito automático', 129)
        add(D(y, m, 12), 'Streaming de vídeo', 'Despesa', 'Assinaturas', 'Streaming', CR, 'Crédito', 55.90)
        add(D(y, m, 12), 'Streaming de música', 'Despesa', 'Assinaturas', 'Música', CR, 'Crédito', 21.90)
        add(D(y, m, 14), 'Armazenamento em nuvem', 'Despesa', 'Assinaturas', 'Nuvem', AZ, 'Crédito', 14.90)
        add(D(y, m, 20), 'Parcela financiamento do carro', 'Despesa', 'Dívidas e Financiamentos', 'Financiamento', CC, 'Débito automático', 1150)
        add(D(y, m, 22), 'Parcela empréstimo pessoal', 'Despesa', 'Dívidas e Financiamentos', 'Empréstimo', CC, 'Débito automático', 580)
        add(D(y, m, 25), 'Parcela renegociação da fatura', 'Despesa', 'Dívidas e Financiamentos', None, CC, 'Boleto', 460)
        # variáveis
        for d in (4, 11, 18, 26):
            card = rnd.choice([AZ, AZ, CC])
            add(D(y, m, min(d + rnd.randint(0, 2), 28)), 'Supermercado', 'Despesa', 'Alimentação', 'Mercado', card,
                'Crédito' if card == AZ else 'Débito', round(rnd.uniform(230, 420) * (1.35 if (y, m) == (2026, 9) else 1), 2))
        for _ in range(2):
            add(D(y, m, rnd.randint(1, 28)), rnd.choice(['Restaurante', 'Delivery de comida', 'Padaria e café']), 'Despesa',
                'Alimentação', rnd.choice(['Restaurantes', 'Delivery']), rnd.choice([CR, AZ]), 'Crédito', round(rnd.uniform(40, 110), 2))
        for _ in range(2):
            add(D(y, m, rnd.randint(1, 28)), 'Combustível', 'Despesa', 'Transporte', 'Combustível', rnd.choice([CR, AZ]), 'Crédito',
                round(rnd.uniform(150, 215), 2))
        for _ in range(3):
            add(D(y, m, rnd.randint(1, 28)), 'Corrida de aplicativo', 'Despesa', 'Transporte', 'App de transporte', CC, 'Pix',
                round(rnd.uniform(18, 48), 2))
        add(D(y, m, rnd.randint(2, 27)), 'Farmácia', 'Despesa', 'Saúde', 'Farmácia', rnd.choice([CR, AZ]), 'Crédito', round(rnd.uniform(40, 170), 2))
        if m % 3 == 0:
            add(D(y, m, rnd.randint(2, 27)), 'Consulta médica', 'Despesa', 'Saúde', 'Consultas', CC, 'Pix', 280)
        for _ in range(2):
            add(D(y, m, rnd.randint(1, 28)), rnd.choice(['Cinema', 'Show', 'Bar com amigos']), 'Despesa', 'Lazer',
                rnd.choice(['Cinema', 'Bares']), rnd.choice([CR, AZ]), 'Crédito', round(rnd.uniform(60, 190), 2))
        if m % 2 == 0:
            add(D(y, m, rnd.randint(1, 28)), 'Roupas', 'Despesa', 'Vestuário', 'Roupas', rnd.choice([CR, AZ]), 'Crédito', round(rnd.uniform(150, 450), 2))
        add(D(y, m, 7), 'Doação mensal', 'Despesa', 'Doações', 'ONG', CC, 'Pix', 100)
        add(D(y, m, 9), 'Ração e petiscos', 'Despesa', 'Família/Pets', 'Ração', AZ, 'Crédito', round(rnd.uniform(190, 250), 2))
        if m in (4, 8):
            add(D(y, m, 16), 'Veterinário', 'Despesa', 'Família/Pets', 'Veterinário', CC, 'Pix', 320)
        add(D(y, m, rnd.randint(2, 27)), 'Despesa diversa', 'Despesa', 'Outros', None, CC, 'Pix', round(rnd.uniform(50, 210), 2))
        # aportes
        add(D(y, m, 6), 'Aporte - reserva de emergência', 'Investimento', 'Investimento', None, CC, 'Transferência', 700, meta='Reserva de emergência')
        add(D(y, m, 6), 'Aporte - viagem', 'Investimento', 'Investimento', None, CC, 'Transferência', 300, meta='Viagem Europa')
        add(D(y, m, 6), 'Aporte - troca de carro', 'Investimento', 'Investimento', None, CC, 'Transferência', 200, meta='Troca de carro')
        if y == 2026:
            add(D(y, m, 6), 'Aporte - MBA', 'Investimento', 'Investimento', None, CC, 'Transferência', 150, meta='MBA')
        add(D(y, m, 6), 'Aporte - aposentadoria (PGBL)', 'Investimento', 'Investimento', None, CC, 'Transferência', 300, meta='Aposentadoria')
        if (y, m) >= (2026, 4):
            add(D(y, m, 6), 'Aporte - quitar empréstimo', 'Investimento', 'Investimento', None, CC, 'Transferência', 150, meta='Quitar empréstimo pessoal')
        if (y, m) >= (2026, 4):
            add(D(y, m, 6), 'Aporte - entrada do apartamento', 'Investimento', 'Investimento', None, CC, 'Transferência', 250, meta='Entrada do apartamento')
    # compras especiais (parceladas / tags)
    L.append([D(2026, 2, 12), 'IPTU 2026 (10 parcelas)', 'Despesa', 'Impostos e Taxas', 'IPTU', CC, 'Boleto', 1, 10, 'Pago', None, None, 1800.0])
    L.append([D(2026, 3, 12), 'Notebook novo (12x)', 'Despesa', 'Outros', None, CR, 'Crédito', 1, 12, 'Pendente', None, None, 3600.0])
    L.append([D(2026, 7, 8), 'Geladeira nova (10x)', 'Despesa', 'Moradia', None, AZ, 'Crédito', 1, 10, 'Pendente', None, None, 2000.0])
    L.append([D(2026, 8, 20), 'Passagens aéreas viagem (4x)', 'Despesa', 'Lazer', 'Viagens', CR, 'Crédito', 1, 4, 'Pendente', '#viagem', None, 1800.0])
    L.append([D(2026, 6, 14), 'Presente de aniversário', 'Despesa', 'Outros', None, AZ, 'Crédito', None, None, 'Pago', '#presente', None, 280.0])
    L.append([D(2026, 9, 18), 'Hospedagem viagem de fim de semana', 'Despesa', 'Lazer', 'Viagens', AZ, 'Crédito', None, None, 'Pago', '#viagem', None, 950.0])
    L.append([D(2026, 10, 18), 'Seguro do carro (anual, em 3x)', 'Despesa', 'Transporte', 'Manutenção', CC, 'Boleto', 1, 3, 'Agendado', None, None, 2700.0])
    L.append([D(2026, 11, 10), 'Aporte extra - 13º (previsto)', 'Investimento', 'Investimento', None, CC, 'Transferência', None, None, 'Agendado', None, 'MBA', 1500.0])
    # pagamentos de fatura (transferência negativa), calculados a partir das compras no cartão
    cards = {CR: (25, 5), AZ: (10, 17)}
    fat = {}
    for row in L:
        d, desc, tipo, cat, sub, conta, forma, pa, pt, status, tag, meta, valor = row
        if tipo == 'Despesa' and conta in cards:
            fech, ven = cards[conta]
            idx = d.year * 12 + d.month + (1 if d.day >= fech else 0)
            n = pt or 1
            for k in range((pa or 1) - 1, n):
                fat[(conta, idx + k - ((pa or 1) - 1))] = fat.get((conta, idx + k - ((pa or 1) - 1)), 0) + valor / n
    for (conta, idx), v in sorted(fat.items(), key=lambda x: x[0][1]):
        fech, ven = cards[conta]
        vi = idx + (0 if ven > fech else 1)
        y, m = divmod(vi - 1, 12)
        m += 1
        due = D(y, m, ven)
        if due <= D(2026, 12, 31) and due >= D(2025, 10, 1):
            L.append([due, 'Pagamento fatura ' + conta, 'Transferência', None, None, CC, 'Débito automático', None, None,
                      'Pago' if due < HOJE else 'Agendado', None, None, -round(v, 2)])
    L.sort(key=lambda r: (r[0], r[1]))
    return L


def fill_lancamentos(ctx):
    ws = ctx.wb['LANÇAMENTOS']
    rows = gerar_lancamentos()
    assert len(rows) < 1400, len(rows)
    for i, row in enumerate(rows):
        r = LR0 + i
        d, desc, tipo, cat, sub, conta, forma, pa, pt, status, tag, meta, valor = row
        vals = [d, desc, tipo, cat, sub, conta, forma, pa, pt, status, tag, meta, 'EXEMPLO', valor]
        for j, v in enumerate(vals):
            if v is not None:
                ws.cell(row=r, column=1 + j, value=v)
    return len(rows)


# ------------------------------------------------------------------ demais abas
def fill_rest(ctx):
    put_range(ctx, 'Rec_Desc', [[x] for x in ['Salário', 'Aluguel recebido', 'Aluguel residencial', 'Condomínio', 'Energia + água',
                                               'Internet fibra', 'Plano de saúde', 'Curso de inglês', 'Streaming de vídeo', 'Streaming de música',
                                               'Armazenamento em nuvem', 'Academia', 'Financiamento do carro', 'Empréstimo pessoal',
                                               'Aporte mensal - reserva', '13º salário', 'IPVA', 'Seguro do carro']])
    rec = [
        ('Receita', 'Salário', 'Conta Corrente Banco Azul', 10800, 'Mensal', 5, D(2026, 1, 1), None),
        ('Receita', 'Aluguel recebido', 'Conta Corrente Banco Azul', 1900, 'Mensal', 10, D(2026, 1, 1), None),
        ('Despesa', 'Moradia', 'Conta Corrente Banco Azul', 2400, 'Mensal', 10, D(2026, 1, 1), None),
        ('Despesa', 'Moradia', 'Conta Corrente Banco Azul', 650, 'Mensal', 10, D(2026, 1, 1), None),
        ('Despesa', 'Moradia', 'Conta Corrente Banco Azul', 260, 'Mensal', 15, D(2026, 1, 1), None),
        ('Despesa', 'Moradia', 'Conta Corrente Banco Azul', 129.9, 'Mensal', 20, D(2026, 1, 1), None),
        ('Despesa', 'Saúde', 'Conta Corrente Banco Azul', 780, 'Mensal', 5, D(2026, 1, 1), None),
        ('Despesa', 'Educação', 'Conta Corrente Banco Azul', 390, 'Mensal', 8, D(2026, 1, 1), D(2026, 12, 31)),
        ('Despesa', 'Assinaturas', 'Cartão Roxinho', 55.9, 'Mensal', 12, D(2026, 1, 1), None),
        ('Despesa', 'Assinaturas', 'Cartão Roxinho', 21.9, 'Mensal', 12, D(2026, 1, 1), None),
        ('Despesa', 'Assinaturas', 'Cartão Azul Gold', 14.9, 'Mensal', 14, D(2026, 1, 1), None),
        ('Despesa', 'Assinaturas', 'Conta Corrente Banco Azul', 129, 'Mensal', 3, D(2026, 1, 1), None),
        ('Despesa', 'Dívidas e Financiamentos', 'Conta Corrente Banco Azul', 1150, 'Mensal', 20, D(2026, 1, 1), D(2029, 3, 31)),
        ('Despesa', 'Dívidas e Financiamentos', 'Conta Corrente Banco Azul', 580, 'Mensal', 22, D(2026, 1, 1), D(2027, 9, 30)),
        ('Investimento', 'Investimento', 'Conta Corrente Banco Azul', 2050, 'Mensal', 6, D(2026, 1, 1), None),
        ('Receita', '13º/Férias/PLR', 'Conta Corrente Banco Azul', 10800, 'Anual', 20, D(2026, 12, 20), None),
        ('Despesa', 'Impostos e Taxas', 'Conta Corrente Banco Azul', 2100, 'Anual', 15, D(2027, 1, 15), None),
        ('Despesa', 'Transporte', 'Conta Corrente Banco Azul', 2700, 'Anual', 18, D(2026, 10, 18), None),
    ]
    put_range(ctx, 'Rec_Tipo', [[r[0]] for r in rec])
    put_range(ctx, 'Rec_Cat', [[r[1]] for r in rec])
    put_range(ctx, 'Rec_Conta', [[r[2]] for r in rec])
    put_range(ctx, 'Rec_Valor', [[r[3]] for r in rec])
    put_range(ctx, 'Rec_Freq', [[r[4]] for r in rec])
    put_range(ctx, 'Rec_Dia', [[r[5]] for r in rec])
    put_range(ctx, 'Rec_Ini', [[r[6]] for r in rec])
    put_range(ctx, 'Rec_Fim', [[r[7]] for r in rec])
    put_range(ctx, 'Rec_Ativo', [['Sim'] for r in rec])
    # orçamento
    plan_d = [3600, 1800, 900, 1100, 450, 700, 300, 400, 2200, 250, 450, 150, 300]
    put_range(ctx, 'Exemplo_OrcDespesas', [[v] * 12 for v in plan_d])
    plan_r = [[10800] * 12, [0] * 11 + [10800], [1000] * 12, [1900] * 12, [80] * 12, [0] * 12, [0] * 12]
    put_range(ctx, 'Exemplo_OrcReceitas', plan_r)
    # metas
    metas = [
        ('Reserva de emergência', 'Reserva', 'Alta', 60000, 40000, D(2025, 9, 1), D(2027, 12, 31), 'Poupança Banco Azul'),
        ('Viagem Europa', 'Viagem', 'Média', 18000, 3000, D(2025, 9, 1), D(2027, 7, 31), 'CDB Banco Verde'),
        ('Troca de carro', 'Carro', 'Média', 15000, 6000, D(2025, 9, 1), D(2028, 12, 31), 'Tesouro Selic'),
        ('Entrada do apartamento', 'Casa', 'Alta', 50000, 12000, D(2025, 9, 1), D(2031, 12, 31), 'Tesouro Selic'),
        ('Quitar empréstimo pessoal', 'Quitar dívida', 'Alta', 6100, 3000, D(2026, 3, 1), D(2027, 9, 30), 'Conta Corrente Banco Azul'),
        ('MBA', 'Estudos', 'Baixa', 4500, 2000, D(2025, 9, 1), D(2027, 6, 30), 'CDB Banco Verde'),
        ('Aposentadoria', 'Aposentadoria', 'Média', 100000, 30000, D(2025, 9, 1), D(2046, 12, 31), 'Previdência PGBL'),
        ('Notebook novo', 'Outros', 'Baixa', 6000, 6000, D(2026, 1, 1), D(2026, 6, 30), 'Conta Corrente Banco Azul'),
    ]
    put_range(ctx, 'Exemplo_Metas', [list(m) for m in metas])
    put_range(ctx, 'Exemplo_Planejador', [[30000], [D(2027, 6, 30)], [5000]])
    # dívidas
    put_range(ctx, 'Exemplo_Dividas', [
        ['Financiamento do carro', 'Financiamento veículo', 45000, 28400, 0.020, 1150, 48, 18, 'Price', 'Em dia'],
        ['Empréstimo pessoal', 'Empréstimo', 10000, 6100, 0.016, 580, 24, 12, 'Price', 'Em dia'],
        ['Renegociação da fatura', 'Cartão/Fatura parcelada', 3200, 2600, 0.035, 460, 12, 4, 'Price', 'Em dia'],
    ])
    ctx.wb['DÍVIDAS E FINANCIAMENTOS'][ctx.refs['Exemplo_SimExtra'][1]] = 600
    ctx.wb['DÍVIDAS E FINANCIAMENTOS'][ctx.refs['Exemplo_AmSel'][1]] = 'Financiamento do carro'
    # investimentos
    def vatual(aport, d, f):
        return round(aport * (1.1375 ** ((HOJE - d).days / 365)) * f, 2)
    inv = [
        ('Tesouro Selic 2029', 'Tesouro Direto', 'Banco Azul', D(2024, 3, 10), 22000, 0.99, 'Diária', D(2029, 3, 1)),
        ('CDB 110% do CDI', 'Renda fixa', 'Banco Verde', D(2025, 2, 15), 15000, 1.04, 'No vencimento', D(2027, 2, 15)),
        ('FII Exemplo Logística', 'FIIs', 'Corretora X', D(2025, 5, 2), 8000, 0.97, 'D+1', None),
        ('Ações - cesta exemplo', 'Ações', 'Corretora X', D(2025, 8, 20), 6000, 0.93, 'D+1', None),
        ('ETF índice exemplo', 'ETFs', 'Corretora X', D(2025, 9, 5), 5000, 1.02, 'D+1', None),
        ('Previdência PGBL', 'Previdência', 'Seguradora Y', D(2023, 1, 10), 30000, 0.96, 'Baixa', None),
        ('Cripto - carteira exemplo', 'Cripto', 'Exchange Z', D(2025, 11, 12), 2000, 1.10, 'D+1', None),
        ('ETF exterior exemplo', 'Exterior', 'Corretora X', D(2025, 6, 18), 4000, 1.06, 'D+1', None),
    ]
    put_range(ctx, 'Exemplo_Invest', [[n, t, i, d, ap, vatual(ap, d, f), lq, v] for n, t, i, d, ap, f, lq, v in inv])
    put_range(ctx, 'Exemplo_InvAlvo', [[0.25], [0.20], [0.10], [0.10], [0.05], [0.02], [0.18], [0.10], [0]])
    # patrimônio
    put_range(ctx, 'Exemplo_PatAtivos', [[0], [520000], [48000], [8000]])
    ctx.wb['PATRIMÔNIO LÍQUIDO'][ctx.refs['Exemplo_PatAno0'][1]] = 618000
    ativos = [672000, 679000, 685000, 693000, 700000, 707000, 714000, 722000, 729000]
    pas = [52000, 51200, 50400, 49300, 48200, 47000, 45800, 44600, 43300]
    put_range(ctx, 'Exemplo_PatHist', [[a, p] for a, p in zip(ativos, pas)])
    # imposto de renda
    put_range(ctx, 'Exemplo_IR_Trib', [['Empresa Exemplo S/A', '00.000.000/0001-00', 102000, 11500],
                                       ['Projetos autônomos (carnê-leão)', '000.000.000-00', 7200, 0]])
    put_range(ctx, 'Exemplo_IR_Isent', [['Rendimento da poupança', 'Banco Azul', 720], ['Aluguel recebido (pessoa física, isento no exemplo)', 'Inquilino', 0]])
    put_range(ctx, 'Exemplo_IR_Excl', [['13º salário', 'Empresa Exemplo S/A', 8500]])
    put_range(ctx, 'Exemplo_IR_Bens', [['Apartamento residencial', '12-01', 520000, 520000], ['Veículo', '21-01', 52000, 48000],
                                       ['Conta corrente e poupança', '06-01', 9000, 10450], ['Investimentos', '04-02', 80000, 106000]])
    put_range(ctx, 'Exemplo_IR_Div', [['Financiamento do veículo', 'Banco Azul', 40900, 28400]])
    ctx.wb['IMPOSTO DE RENDA'][ctx.refs['Exemplo_IR_Ded2'][1]] = 0
    # simulador "E se?"
    put_range(ctx, 'Exemplo_Simulador', [[300], ['Lazer'], [500], [0.10]])
    # backup
    ctx.wb['BACKUP'][ctx.refs['Exemplo_Backup'][1]] = D(2026, 9, 28)
    # dashboard: mês completo
    ws = ctx.wb['DASHBOARD']
    ws['C6'] = 'Setembro'
    ws['E6'] = 2026


def fill_sample(ctx):
    fill_config(ctx)
    n = fill_lancamentos(ctx)
    fill_rest(ctx)
    return n
