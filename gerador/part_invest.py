"""INVESTIMENTOS, PATRIMÔNIO LÍQUIDO e IMPOSTO DE RENDA."""
from openpyxl.chart import LineChart, Reference, Series
from common import *
from part_base import TIPOS_INV, MESES_AB, N_CATR

I0, N_INV = 10, 30
I1 = I0 + N_INV - 1


def build_invest(ctx):
    sh = Sh(ctx, 'INVESTIMENTOS')
    ws = sh.ws
    S = 'INVESTIMENTOS'
    sh.header('INVESTIMENTOS', 'Cadastre seus ativos (células azuis). A planilha calcula rentabilidade, % da carteira, compara com '
              'CDI e inflação (de CONFIGURAÇÕES) e indica rebalanceamento. ⚠ Apenas controle pessoal — não é recomendação de investimento.',
              ncols=17)
    sh.widths({'B': 28, 'C': 16, 'D': 18, 'E': 15, 'F': 15, 'G': 15, 'H': 14, 'I': 13, 'J': 14, 'K': 14, 'L': 14, 'M': 14,
               'N': 14, 'O': 12, 'P': 12, 'Q': 22})
    sh.put('B5', '⬇ digite nas colunas azuis (B a I)', 'note', color=COR['azul'], bold=True)
    sh.put('B6', '1. Carteira de ativos', 'sec')
    sh.put('B7', 'Aviso: esta aba serve apenas para organizar e acompanhar seus investimentos. Não constitui recomendação '
           'de compra, venda ou alocação.', 'note')
    heads = ['Ativo', 'Tipo', 'Instituição', 'Data da aplicação', 'Valor aportado', 'Valor atual', 'Liquidez', 'Vencimento',
             'Rentabilidade R$', 'Rentab. %', '% da carteira', 'Dias', 'CDI no período', 'Inflação no período',
             'Rentab. − CDI (p.p.)', 'Situação']
    for i, h in enumerate(heads):
        sh.put(f'{CL(2 + i)}9', h, 'hdr')
    ws.row_dimensions[9].height = 46
    for r in range(I0, I1 + 1):
        for L, fm in zip('BCDEFGHI', [None, None, None, DATA, BRLZ, BRLZ, None, DATA]):
            sh.put(f'{L}{r}', None, 'in', fm)
        f = {
            'J': f'=IF($B{r}="",0,N($G{r})-N($F{r}))',
            'K': f'=IF(N($F{r})=0,0,J{r}/$F{r})',
            'L': f'=IF($B{r}="",0,IFERROR($G{r}/Inv_TotAtual,0))',
            'M': f'=IF(OR($B{r}="",NOT(ISNUMBER($E{r}))),0,MAX(1,Cfg_Hoje-$E{r}))',
            'N': f'=IF(M{r}=0,0,(1+Cfg_CDI)^(M{r}/365)-1)',
            'O': f'=IF(M{r}=0,0,(1+Cfg_Infl)^(M{r}/365)-1)',
            'P': f'=IF(M{r}=0,0,(K{r}-N{r})*100)',
            'Q': f'=IF($B{r}="","",IF(M{r}=0,"—",IF(K{r}>=N{r},"✔ Acima do CDI",IF(K{r}>=O{r},"Acima da inflação","⚠ Abaixo da inflação"))))',
        }
        fm = {'J': BRL, 'K': PCT, 'L': PCT, 'M': INT, 'N': PCT, 'O': PCT, 'P': '0.0;-0.0;0.0'}
        for L, fx in f.items():
            sh.put(f'{L}{r}', fx, 'calc', fm.get(L))
    t = I1 + 1
    sh.put(f'B{t}', 'TOTAL DA CARTEIRA', 'calc', bold=True)
    for L in 'FGJ':
        sh.put(f'{L}{t}', f'=SUM({L}{I0}:{L}{I1})', 'calc', BRL, bold=True)
    sh.put(f'K{t}', f'=IF(F{t}=0,0,J{t}/F{t})', 'calc', PCT, bold=True)
    sh.put(f'L{t}', f'=SUM(L{I0}:L{I1})', 'calc', PCT, bold=True)
    sh.put(f'N{t}', f'=IF(F{t}=0,0,SUMPRODUCT(N{I0}:N{I1},F{I0}:F{I1})/F{t})', 'calc', PCT, bold=True)
    sh.put(f'O{t}', f'=IF(F{t}=0,0,SUMPRODUCT(O{I0}:O{I1},F{I0}:F{I1})/F{t})', 'calc', PCT, bold=True)
    sh.put(f'P{t}', f'=(K{t}-N{t})*100', 'calc', '0.0;-0.0;0.0', bold=True)
    sh.put(f'Q{t}', f'=IF(F{t}=0,"",IF(K{t}>=N{t},"✔ Carteira acima do CDI",IF(K{t}>=O{t},"Acima da inflação","⚠ Abaixo da inflação")))', 'calc', bold=True)
    sh.put(f'B{t + 1}', 'CDI e inflação "no período" = taxa anual de CONFIGURAÇÕES capitalizada pelos dias desde a aplicação '
           '(aproximação para comparação; renda variável oscila).', 'note')
    ctx.name('Inv_TotAtual', S, f'G{t}')
    ctx.name('Inv_TotAport', S, f'F{t}')
    ctx.name('Inv_Rent', S, f'K{t}')
    for L, n in zip('BCDEFGHI', ['Nome', 'Tipo', 'Inst', 'Data', 'Aport', 'Atual', 'Liq', 'Venc']):
        ctx.name(f'Inv_{n}', S, f'{L}{I0}:{L}{I1}')
    ctx.name('Exemplo_Invest', S, f'B{I0}:I{I1}')
    ctx.clear_ranges.append((S, f'B{I0}:I{I1}', 'Investimentos de exemplo'))
    sh.dv(f'C{I0}:C{I1}', '=Lista_TipoInv')
    sh.dv(f'H{I0}:H{I1}', '=Lista_Liquidez')
    sh.cf_formula(f'Q{I0}:Q{t}', f'LEFT(Q{I0},1)="✔"', COR['verde_l'], COR['verde'], True)
    sh.cf_formula(f'Q{I0}:Q{t}', f'LEFT(Q{I0},1)="⚠"', COR['vermelho_l'], COR['vermelho'], True)
    sh.cf_formula(f'J{I0}:K{t}', f'J{I0}<0', None, COR['vermelho'])
    sh.cf_formula(f'J{I0}:K{t}', f'J{I0}>0', None, COR['verde'])
    # ---- alocação atual x alvo
    a = t + 3
    sh.put(f'B{a}', '2. Alocação atual x alocação-alvo (digite o alvo em %)', 'sec')
    sh.put(f'F{a}', 'Tolerância (p.p.):', 'lbl', h='right')
    sh.put(f'G{a}', None, 'in', '0.0')
    ws[f'G{a}'] = 5
    ctx.name('Inv_Tol', S, f'G{a}')
    for i, h in enumerate(['Tipo de ativo', 'Valor atual', '% atual', '% alvo', 'Diferença (p.p.)', 'Valor alvo', 'Ajuste necessário',
                           'Indicação']):
        sh.put(f'{CL(2 + i)}{a + 1}', h, 'hdr')
    k0 = a + 2
    for i, tp in enumerate(TIPOS_INV):
        r = k0 + i
        sh.put(f'B{r}', tp, 'calc')
        sh.put(f'C{r}', f'=SUMIF(Inv_Tipo,B{r},Inv_Atual)', 'calc', BRL)
        sh.put(f'D{r}', f'=IFERROR(C{r}/Inv_TotAtual,0)', 'calc', PCT)
        sh.put(f'E{r}', None, 'in', PCT)
        sh.put(f'F{r}', f'=IF(E{r}="","",(D{r}-E{r})*100)', 'calc', '0.0;-0.0;0.0')
        sh.put(f'G{r}', f'=IF(E{r}="",0,E{r}*Inv_TotAtual)', 'calc', BRL)
        sh.put(f'H{r}', f'=IF(E{r}="",0,G{r}-C{r})', 'calc', BRLZ)
        sh.put(f'I{r}', f'=IF(COUNT($E${k0}:$E${k0 + 8})=0,"defina o alvo",IF(ABS(N(F{r}))<=Inv_Tol,"✔ Ok",'
               f'IF(F{r}>0,"Reduzir R$ "&FIXED(-H{r},0),"Aportar R$ "&FIXED(H{r},0))))', 'calc')
    k1 = k0 + 8
    r = k1 + 1
    sh.put(f'B{r}', 'Total', 'calc', bold=True)
    sh.put(f'C{r}', f'=SUM(C{k0}:C{k1})', 'calc', BRL, bold=True)
    sh.put(f'D{r}', f'=SUM(D{k0}:D{k1})', 'calc', PCT, bold=True)
    sh.put(f'E{r}', f'=SUM(E{k0}:E{k1})', 'calc', PCT, bold=True)
    sh.put(f'I{r}', f'=IF(COUNT(E{k0}:E{k1})=0,"",IF(ROUND(E{r},4)=1,"✔ Alvo soma 100%","⚠ O alvo deve somar 100%"))', 'calc', bold=True)
    ctx.name('Inv_TipoNome', S, f'B{k0}:B{k1}')
    ctx.name('Inv_TipoVal', S, f'C{k0}:C{k1}')
    ctx.name('Exemplo_InvAlvo', S, f'E{k0}:E{k1}')
    ctx.clear_ranges.append((S, f'E{k0}:E{k1}', 'Alocação-alvo de exemplo'))
    sh.cf_formula(f'I{k0}:I{r}', f'OR(LEFT(I{k0},7)="Reduzir",LEFT(I{k0},7)="Aportar")', COR['amarelo_l'], '9A5B00', True)
    sh.cf_formula(f'I{k0}:I{r}', f'LEFT(I{k0},1)="✔"', COR['verde_l'], COR['verde'], True)
    sh.cf_formula(f'I{k0}:I{r}', f'LEFT(I{k0},1)="⚠"', COR['vermelho_l'], COR['vermelho'], True)
    # ---- aportes mensais e evolução (ano selecionado)
    e = r + 3
    sh.put(f'B{e}', '3. Aportes mensais e evolução da carteira (ano selecionado no DASHBOARD)', 'sec')
    sh.put(f'B{e + 1}', 'Item', 'hdr')
    for k, m in enumerate(MESES_AB):
        sh.put(f'{CL(3 + k)}{e + 1}', m, 'hdr')
    labels = ['Aportes lançados no mês (LANÇAMENTOS)', 'Aportes acumulados no ano', 'Carteira: valor aportado acumulado (por data)',
              'Se rendesse o CDI (aproximação)', 'Se fosse corrigida pela inflação (aproximação)']
    for i, lb in enumerate(labels):
        sh.put(f'B{e + 2 + i}', lb, 'calc', bold=(i == 2))
    for k in range(12):
        L = CL(3 + k)
        P = CL(2 + k)
        idx = f'(Sel_Ano*12+{k + 1})'
        sh.put(f'{L}{e + 2}', f'={SM("Investimento", idx)}', 'calc', BRL)
        sh.put(f'{L}{e + 3}', f'=SUM($C{e + 2}:{L}{e + 2})', 'calc', BRL)
        sh.put(f'{L}{e + 4}', f'=SUMIFS(Inv_Aport,Inv_Data,"<="&(DATE(Sel_Ano,{k + 2},1)-1))', 'calc', BRL)
        if k == 0:
            sh.put(f'{L}{e + 5}', f'={L}{e + 4}', 'calc', BRL)
            sh.put(f'{L}{e + 6}', f'={L}{e + 4}', 'calc', BRL)
        else:
            sh.put(f'{L}{e + 5}', f'={P}{e + 5}*(1+Cfg_CDI)^(1/12)+({L}{e + 4}-{P}{e + 4})', 'calc', BRL)
            sh.put(f'{L}{e + 6}', f'={P}{e + 6}*(1+Cfg_Infl)^(1/12)+({L}{e + 4}-{P}{e + 4})', 'calc', BRL)
    sh.put(f'O{e + 1}', 'Total', 'hdr')
    sh.put(f'O{e + 2}', f'=SUM(C{e + 2}:N{e + 2})', 'calc', BRL, bold=True)
    sh.finish(freeze='C10')
    return sh


# ======================================================================
def build_patrimonio(ctx):
    sh = Sh(ctx, 'PATRIMÔNIO LÍQUIDO')
    ws = sh.ws
    S = 'PATRIMÔNIO LÍQUIDO'
    sh.header('PATRIMÔNIO LÍQUIDO', 'Ativos (o que você tem) menos Passivos (o que você deve) = patrimônio líquido. Contas, '
              'investimentos, dívidas e faturas vêm automaticamente; bens você digita. No fim de cada mês, registre um "retrato" no histórico.',
              ncols=9)
    sh.widths({'B': 40, 'C': 18, 'D': 18, 'E': 18, 'F': 16, 'G': 14, 'H': 20, 'I': 12})
    sh.put('B6', '1. Situação atual', 'sec')
    sh.put('B7', 'ATIVOS', 'lblb', color=COR['verde'])
    ativos = [('Saldo em contas (CONFIGURAÇÕES)', '=SUM(Contas_SaldoAtual)', None),
              ('Investimentos (INVESTIMENTOS)', '=Inv_TotAtual', None),
              ('Poupança / reservas fora do cadastro', None, 'in'), ('Imóveis', None, 'in'), ('Veículos', None, 'in'),
              ('Outros bens', None, 'in')]
    for i, (lb, fx, k) in enumerate(ativos):
        r = 8 + i
        sh.put(f'B{r}', lb, 'lbl')
        sh.put(f'C{r}', fx, k or 'calc', BRLZ)
    sh.put('B14', 'Total de ativos', 'calc', bold=True)
    sh.put('C14', '=SUM(C8:C13)', 'calc', BRLZ, bold=True)
    sh.put('B16', 'PASSIVOS', 'lblb', color=COR['vermelho'])
    pas = [('Dívidas e financiamentos (DÍVIDAS)', '=Div_Total', None), ('Faturas de cartão em aberto (CARTÕES)', '=Card_FatAntTot+Card_FatAtualTot', None),
           ('Outros passivos', None, 'in')]
    for i, (lb, fx, k) in enumerate(pas):
        r = 17 + i
        sh.put(f'B{r}', lb, 'lbl')
        sh.put(f'C{r}', fx, k or 'calc', BRLZ)
    sh.put('B20', 'Total de passivos', 'calc', bold=True)
    sh.put('C20', '=SUM(C17:C19)', 'calc', BRLZ, bold=True)
    sh.put('B22', 'PATRIMÔNIO LÍQUIDO ATUAL', 'calc', bold=True, size=13)
    sh.put('C22', '=C14-C20', 'calc', BRLZ, bold=True, size=13)
    sh.cf_formula('C22', 'C22<0', COR['vermelho_l'], COR['vermelho'], True)
    sh.cf_formula('C22', 'C22>=0', COR['verde_l'], COR['verde'], True)
    sh.put('D8', 'Se mantiver parte do dinheiro numa "Conta investimento" E também em INVESTIMENTOS, lance só uma vez para não duplicar.', 'note')
    ctx.name('Pat_Ativos', S, 'C14')
    ctx.name('Pat_Passivos', S, 'C20')
    ctx.name('Pat_Liquido', S, 'C22')
    ctx.name('Exemplo_PatAtivos', S, 'C10:C13')
    ctx.name('Exemplo_PatPassivo', S, 'C19')
    ctx.clear_ranges.append((S, 'C10:C13', 'Bens de exemplo'))
    # ---- histórico
    h = 25
    sh.put(f'B{h}', '2. Histórico mensal (retrato manual no fim de cada mês) — ano selecionado no DASHBOARD', 'sec')
    sh.put(f'B{h + 1}', 'Patrimônio líquido em 31/12 do ano anterior:', 'lblb')
    sh.put(f'C{h + 1}', None, 'in', BRLZ)
    ctx.name('Pat_Ano0', S, f'C{h + 1}')
    ctx.name('Exemplo_PatAno0', S, f'C{h + 1}')
    sh.put(f'D{h + 1}', '← digite aqui (ponto de partida do ano)', 'note')
    for i, t in enumerate(['Mês', 'Total de ativos', 'Total de passivos', 'Patrimônio líquido', 'Variação R$', 'Variação %',
                           'Série do gráfico (retrato ou estimativa)']):
        sh.put(f'{CL(2 + i)}{h + 2}', t, 'hdr')
    ws.row_dimensions[h + 2].height = 46
    r0 = h + 3
    for k in range(12):
        r = r0 + k
        sh.put(f'B{r}', MESES_AB[k], 'calc', bold=True)
        sh.put(f'C{r}', None, 'in', BRLZ)
        sh.put(f'D{r}', None, 'in', BRLZ)
        sh.put(f'E{r}', f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),C{r}-D{r},"")', 'calc', BRLZ)
        prev = 'Pat_Ano0' if k == 0 else f'E{r - 1}'
        sh.put(f'F{r}', f'=IF(AND(ISNUMBER(E{r}),ISNUMBER({prev})),E{r}-{prev},"")', 'calc', BRLZ)
        sh.put(f'G{r}', f'=IF(AND(ISNUMBER(F{r}),N({prev})<>0),F{r}/ABS({prev}),"")', 'calc', PCT)
        prev_s = 'N(Pat_Ano0)' if k == 0 else f'H{r - 1}'
        sh.put(f'H{r}', f'=IF(ISNUMBER(E{r}),E{r},{prev_s}+INDEX(Flux_Saldo,{k + 1}))', 'calc', BRLZ)
    ctx.name('Exemplo_PatHist', S, f'C{r0}:D{r0 + 11}')
    ctx.clear_ranges.append((S, f'C{r0}:D{r0 + 11}', 'Histórico de patrimônio de exemplo'))
    ctx.name('Pat_Serie', S, f'H{r0}:H{r0 + 11}')
    ctx.name('Pat_Meses', S, f'B{r0}:B{r0 + 11}')
    sh.cf_formula(f'F{r0}:G{r0 + 11}', f'AND(ISNUMBER(F{r0}),F{r0}<0)', None, COR['vermelho'], True)
    sh.cf_formula(f'F{r0}:G{r0 + 11}', f'AND(ISNUMBER(F{r0}),F{r0}>0)', None, COR['verde'], True)
    sh.put(f'B{r0 + 12}', 'Dica: no último dia do mês copie "Total de ativos" e "Total de passivos" (seção 1) e cole como VALOR nas '
           'colunas azuis. Meses sem retrato são estimados (patrimônio anterior + saldo do mês).', 'note')
    ch = LineChart()
    ch.title = 'Evolução do patrimônio líquido no ano'
    ch.height, ch.width = 8, 18
    ch.add_data(Reference(ws, min_col=8, min_row=h + 2, max_row=r0 + 11), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=r0, max_row=r0 + 11))
    ch.series[0].graphicalProperties.line.solidFill = COR['azul']
    ch.series[0].graphicalProperties.line.width = 28000
    ch.series[0].smooth = False
    ch.legend.position = 'b'
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.visible_cells_only = False
    ws.add_chart(ch, f'B{r0 + 14}')
    sh.finish(freeze='A6')
    return sh


# ======================================================================
def build_ir(ctx):
    sh = Sh(ctx, 'IMPOSTO DE RENDA')
    ws = sh.ws
    S = 'IMPOSTO DE RENDA'
    sh.header('IMPOSTO DE RENDA (apoio)', 'Organize os dados para a declaração do ano selecionado. ⚠ Apenas organização: não substitui '
              'contador nem o programa oficial da Receita Federal. Valores de limites são referência (edite conforme o ano-calendário).',
              ncols=7)
    sh.widths({'B': 40, 'C': 22, 'D': 18, 'E': 18, 'F': 18, 'G': 30})
    sh.put('B5', '⬇ digite nas células azuis', 'note', color=COR['azul'], bold=True)
    y = [6]

    def tabela(titulo, heads, n, fmts, nome, note=None, total_col=None):
        r = y[0]
        sh.put(f'B{r}', titulo, 'sec')
        if note:
            sh.put(f'D{r}', note, 'note')
        for i, h in enumerate(heads):
            sh.put(f'{CL(2 + i)}{r + 1}', h, 'hdr')
        r0 = r + 2
        for k in range(n):
            for i, fm in enumerate(fmts):
                sh.put(f'{CL(2 + i)}{r0 + k}', None, 'in', fm)
        r1 = r0 + n - 1
        if total_col:
            sh.put(f'B{r1 + 1}', 'Total', 'calc', bold=True)
            for L in total_col:
                sh.put(f'{L}{r1 + 1}', f'=SUM({L}{r0}:{L}{r1})', 'calc', BRLZ, bold=True)
            ctx.name(f'IR_{nome}', S, f'{total_col[0]}{r1 + 1}')
            if len(total_col) > 1:
                ctx.name(f'IR_{nome}2', S, f'{total_col[1]}{r1 + 1}')
        ctx.name(f'Exemplo_IR_{nome}', S, f'B{r0}:{CL(1 + len(heads))}{r1}')
        ctx.clear_ranges.append((S, f'B{r0}:{CL(1 + len(heads))}{r1}', titulo))
        y[0] = r1 + 3
        return r0, r1

    tabela('1. Rendimentos tributáveis (informes de rendimentos)', ['Fonte pagadora', 'CNPJ/CPF', 'Rendimento anual', 'IRRF retido'],
           6, [None, None, BRLZ, BRLZ], 'Trib', total_col='DE')
    tabela('2. Rendimentos isentos e não tributáveis', ['Descrição', 'Fonte', 'Valor anual'], 6, [None, None, BRLZ], 'Isent',
           total_col='D')
    tabela('3. Tributação exclusiva / definitiva (ex.: 13º, aplicações)', ['Descrição', 'Fonte', 'Valor anual'], 5,
           [None, None, BRLZ], 'Excl', total_col='D')
    # deduções
    r = y[0]
    sh.put(f'B{r}', '4. Despesas dedutíveis', 'sec')
    for i, h in enumerate(['Item', 'Lançado na planilha (ano)', 'Ajuste manual (+/−)', 'Limite / parâmetro', 'Dedução considerada']):
        sh.put(f'{CL(2 + i)}{r + 1}', h, 'hdr')
    ws.row_dimensions[r + 1].height = 32
    d = r + 2
    sh.put(f'B{d}', 'Saúde (sem limite)', 'lbl')
    sh.put(f'C{d}', '=SUMIFS(L_VParc,L_Tipo,"Despesa",L_Cat,"Saúde",L_Ano,Sel_Ano)', 'calc', BRLZ)
    sh.put(f'D{d}', None, 'in', BRLZ)
    sh.put(f'E{d}', 'sem limite', 'calc', h='center')
    sh.put(f'F{d}', f'=MAX(0,C{d}+N(D{d}))', 'calc', BRLZ)
    sh.put(f'B{d + 1}', 'Educação (limite por pessoa)', 'lbl')
    sh.put(f'C{d + 1}', '=SUMIFS(L_VParc,L_Tipo,"Despesa",L_Cat,"Educação",L_Ano,Sel_Ano)', 'calc', BRLZ)
    sh.put(f'D{d + 1}', None, 'in', BRLZ)
    sh.put(f'E{d + 1}', None, 'in', BRLZ)
    ws[f'E{d + 1}'] = 3561.50
    sh.put(f'F{d + 1}', f'=MIN(MAX(0,C{d + 1}+N(D{d + 1})),E{d + 1}*(1+N(C{d + 2})))', 'calc', BRLZ)
    sh.put(f'B{d + 2}', 'Dependentes: nº na coluna C × dedução por dependente', 'lbl')
    sh.put(f'C{d + 2}', None, 'in', INT)
    sh.put(f'E{d + 2}', None, 'in', BRLZ)
    ws[f'E{d + 2}'] = 2275.08
    sh.put(f'F{d + 2}', f'=N(C{d + 2})*E{d + 2}', 'calc', BRLZ)
    sh.put(f'B{d + 3}', 'Previdência privada (PGBL, limite 12% da renda tributável)', 'lbl')
    sh.put(f'C{d + 3}', None, 'in', BRLZ)
    sh.put(f'D{d + 3}', None, 'in', BRLZ)
    sh.put(f'E{d + 3}', f'=0.12*IR_Trib', 'calc', BRLZ)
    sh.put(f'F{d + 3}', f'=MIN(MAX(0,N(C{d + 3})+N(D{d + 3})),E{d + 3})', 'calc', BRLZ)
    sh.put(f'B{d + 4}', 'Total de deduções', 'calc', bold=True)
    sh.put(f'F{d + 4}', f'=SUM(F{d}:F{d + 3})', 'calc', BRLZ, bold=True)
    ctx.name('IR_Deducoes', S, f'F{d + 4}')
    ctx.name('Exemplo_IR_Ded', S, f'D{d}:D{d + 3}')
    ctx.name('Exemplo_IR_Ded2', S, f'C{d + 2}')
    ctx.name('Exemplo_IR_Ded3', S, f'C{d + 3}')
    sh.put(f'B{d + 5}', 'Educação: limite anual por pessoa (titular + dependentes) e dedução por dependente são valores de referência '
           'editáveis. Confira os valores oficiais do ano-calendário.', 'note')
    y[0] = d + 8
    tabela('5. Bens e direitos', ['Descrição', 'Código', 'Valor em 31/12 do ano anterior', 'Valor em 31/12 do ano'], 10,
           [None, None, BRLZ, BRLZ], 'Bens', total_col='DE')
    tabela('6. Dívidas e ônus reais', ['Descrição', 'Credor', 'Valor em 31/12 do ano anterior', 'Valor em 31/12 do ano'], 5,
           [None, None, BRLZ, BRLZ], 'Div', total_col='DE')
    # checklist
    r = y[0]
    sh.put(f'B{r}', '7. Checklist da declaração', 'sec')
    for i, h in enumerate(['Item', 'Situação']):
        sh.put(f'{CL(2 + i)}{r + 1}', h, 'hdr')
    itens = ['Informes de rendimentos de todas as fontes pagadoras', 'Informes de bancos e corretoras (saldos e rendimentos)',
             'Recibos de despesas médicas e plano de saúde', 'Comprovantes de educação (próprio e dependentes)',
             'Dados de dependentes (CPF e data de nascimento)', 'Comprovantes de previdência privada (PGBL)',
             'Bens e direitos: documentos de imóveis e veículos', 'Dívidas e financiamentos com saldo em 31/12',
             'Aluguéis recebidos / pagos (carnê-leão, se houver)', 'Ganhos de capital e vendas de ativos no ano',
             'Doações e heranças recebidas ou feitas', 'Conferir tudo no programa oficial e guardar o recibo']
    c0 = r + 2
    for i, it in enumerate(itens):
        sh.put(f'B{c0 + i}', it, 'lbl')
        sh.put(f'C{c0 + i}', None, 'in')
        ws[f'C{c0 + i}'] = 'Pendente'
    sh.dv(f'C{c0}:C{c0 + len(itens) - 1}', '=Lista_StatusIR')
    ctx.name('Exemplo_IR_Check', S, f'C{c0}:C{c0 + len(itens) - 1}')
    sh.cf_formula(f'C{c0}:C{c0 + len(itens) - 1}', f'C{c0}="Feito"', COR['verde_l'], COR['verde'], True)
    sh.cf_formula(f'C{c0}:C{c0 + len(itens) - 1}', f'C{c0}="Pendente"', COR['amarelo_l'], '9A5B00', True)
    cn = c0 + len(itens)
    sh.put(f'B{cn}', 'Progresso do checklist', 'calc', bold=True)
    sh.put(f'C{cn}', f'=IFERROR(COUNTIF(C{c0}:C{cn - 1},"Feito")/(ROWS(C{c0}:C{cn - 1})-COUNTIF(C{c0}:C{cn - 1},"Não se aplica")),0)', 'calc', PCT, bold=True)
    # resumo e conferência com lançamentos
    z = cn + 3
    sh.put(f'B{z}', '8. Resumo para a declaração', 'sec')
    resumo = [('Rendimentos tributáveis', '=IR_Trib', BRLZ), ('IRRF retido na fonte', '=IR_Trib2', BRLZ),
              ('Rendimentos isentos e não tributáveis', '=IR_Isent', BRLZ), ('Tributação exclusiva/definitiva', '=IR_Excl', BRLZ),
              ('Deduções legais (item 4)', '=IR_Deducoes', BRLZ),
              ('Base de cálculo estimada (tributáveis − deduções)', '=MAX(0,IR_Trib-IR_Deducoes)', BRLZ),
              ('Bens e direitos em 31/12 do ano', '=IR_Bens2', BRLZ), ('Dívidas e ônus em 31/12 do ano', '=IR_Div2', BRLZ),
              ('Variação patrimonial declarada (bens − dívidas)', '=IR_Bens2-IR_Div2-(IR_Bens-IR_Div)', BRLZ)]
    for i, (lb, fx, fm) in enumerate(resumo):
        sh.put(f'B{z + 1 + i}', lb, 'calc')
        sh.put(f'C{z + 1 + i}', fx, 'calc', fm, bold=True)
    q = z + len(resumo) + 3
    sh.put(f'B{q}', '9. Conferência: receitas lançadas no ano (LANÇAMENTOS) por categoria', 'sec')
    sh.put(f'B{q + 1}', 'Categoria de receita', 'hdr')
    sh.put(f'C{q + 1}', 'Total lançado no ano', 'hdr')
    for k in range(1, N_CATR + 1):
        sh.put(f'B{q + 1 + k}', f'=IF(INDEX(CatR_Nome,{k})="","",INDEX(CatR_Nome,{k}))', 'calc')
        sh.put(f'C{q + 1 + k}', f'=IF(B{q + 1 + k}="",0,SUMIFS(L_VParc,L_Tipo,"Receita",L_Cat,B{q + 1 + k},L_Ano,Sel_Ano))', 'calc', BRL)
    sh.put(f'B{q + 2 + N_CATR}', 'Use para conferir com os informes. Os valores oficiais são os dos informes de rendimentos.', 'note')
    sh.finish(freeze='A6')
    return sh
